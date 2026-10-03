"""Governed unauthorized-debit clearing for SentinelOps Funds Custody.

A sealed transaction or physical debit row is the unit of authorization.
Reconciliation is read-only and deliberately exposes every ambiguity. Checker
approval runs a gated, row-locked, revalidated, atomic Oracle mutation.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
import threading
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Callable, Iterable, Literal
from uuid import uuid4
from zoneinfo import ZoneInfo

import psycopg
from psycopg.rows import dict_row
from pydantic import BaseModel, Field, ValidationError, field_validator, model_validator

from app.config.settings import settings
from app.utils.audit import audit_logger
from app.utils.logging import get_logger


logger = get_logger(__name__)
CLEARING_MIGRATION = (
    "2026_07_add_nexus_unauthorized_clearing.sql followed by "
    "2026_08_refine_nexus_unauthorized_debit_clearing.sql followed by "
    "2026_09_add_nexus_clearing_rrn_lookup.sql followed by "
    "2026_10_refine_nexus_clearing_custody.sql followed by "
    "2026_11_refine_nexus_clearing_workflow.sql followed by "
    "2026_12_repair_nexus_clearing_source_contract.sql followed by "
    "2026_13_refine_nexus_clearing_tranches.sql followed by "
    "2026_18_add_nexus_clearing_queue_custody.sql"
)
MONEY_QUANTUM = Decimal("0.01")
SAFE_STATES = {"READY", "READY_WITH_RESIDUAL", "SAFE_EXTRA_QUEUE"}
IDENTIFIER_RE = re.compile(r"^[A-Z][A-Z0-9_$#]*$")
CLEARING_TRANCHE_SIZE = 20
PHYSICAL_MUTATION_KEYS = {
    "ACNTBAL_AC_UNAUTH_DB_SUM": "unauth_db_sum",
    "ACNTBAL_AC_DB_QUEUE_AMT": "debit_queue_amount",
}
QUEUE_LINKED_MODES = {"QUEUE_EXACT", "QUEUE_PARTIAL", "QUEUE_ROW_ONLY"}


def _money(value: Any) -> Decimal:
    try:
        return Decimal(str(value).replace(",", "").strip()).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
    except (InvalidOperation, AttributeError, ValueError) as exc:
        raise ValueError(f"Invalid monetary value: {value!r}") from exc


def _plain(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return {"type": "bytes", "value": base64.b64encode(value).decode("ascii")}
    if isinstance(value, dict):
        return {str(key): _plain(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_plain(item) for item in value]
    return value


def _json(value: Any) -> str:
    return json.dumps(_plain(value), sort_keys=True, separators=(",", ":"))


def _normalized_text(value: Any) -> str:
    return str(value or "").strip()


def _normalized_account(value: Any) -> str:
    text = _normalized_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if not text or not re.fullmatch(r"[A-Za-z0-9_-]+", text):
        raise ValueError(f"Invalid account identifier: {value!r}")
    return text


def _normalized_reference(value: Any, label: str) -> str:
    text = _normalized_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if not text or len(text) > 128:
        raise ValueError(f"{label} is missing or invalid.")
    return text


def _normalized_rrn_lookup(value: Any) -> str:
    rrn = _normalized_reference(value, "RRN")
    if not re.fullmatch(r"[A-Za-z0-9_-]+", rrn):
        raise ValueError("RRN lookup accepts letters, numbers, hyphens, and underscores only.")
    return rrn


def _currency(value: Any) -> str:
    return _normalized_text(value).upper()


def _select_execution_evidence_rows(value: Any, row_ids: set[str]) -> Any:
    """Keep the source shape while narrowing Oracle evidence to approved rows."""
    if isinstance(value, dict):
        selected: dict[str, Any] = {}
        for key, row in value.items():
            candidate = row.get("row_id") if isinstance(row, dict) else key
            if str(candidate or key) in row_ids:
                selected[key] = row
        return selected
    if isinstance(value, list):
        return [
            row
            for row in value
            if isinstance(row, dict) and str(row.get("row_id") or "") in row_ids
        ]
    return {}


def _execution_item_evidence(
    payload: dict[str, Any],
    evidence: dict[str, Any],
) -> dict[str, dict[str, Any]]:
    """Project the execution-wide Oracle snapshot onto each sealed transaction."""
    before = evidence.get("before") or {}
    after = evidence.get("after") or {}
    mutations = list(evidence.get("balance_mutations") or [])
    queue_deletions = list(evidence.get("queue_deletions") or [])
    projected: dict[str, dict[str, Any]] = {}

    for transaction in payload.get("transactions") or []:
        fingerprint = _normalized_text(transaction.get("fingerprint"))
        if not fingerprint:
            continue
        account = _normalized_text(transaction.get("from_account"))
        currency = _currency(transaction.get("currency"))
        item_mutations = [
            mutation
            for mutation in mutations
            if _normalized_text(mutation.get("external_account")) == account
            and _currency(mutation.get("currency")) == currency
        ]
        item_deletions = [
            deletion
            for deletion in queue_deletions
            if _normalized_text(deletion.get("fingerprint")) == fingerprint
        ]
        balance_row_ids = {
            _normalized_text(mutation.get("row_id"))
            for mutation in item_mutations
            if mutation.get("row_id")
        }
        queue_row_ids = {
            _normalized_text(deletion.get("row_id"))
            for deletion in item_deletions
            if deletion.get("row_id")
        }
        projected[fingerprint] = {
            "before": {
                "balance_rows": _select_execution_evidence_rows(
                    before.get("balance_rows"), balance_row_ids
                ),
                "queue_rows": _select_execution_evidence_rows(
                    before.get("queue_rows"), queue_row_ids
                ),
            },
            "after": {
                "balance_mutations": item_mutations,
                "deleted_queue_row_ids": [
                    row_id
                    for row_id in after.get("deleted_queue_row_ids") or []
                    if _normalized_text(row_id) in queue_row_ids
                ],
            },
        }
    return projected


def _queue_account_currency(queue: dict[str, Any]) -> str:
    """Return the currency that owns the physical ACNTBAL row."""
    return _currency(queue.get("account_currency") or queue.get("currency"))


def _is_explicit_credit_queue(queue: dict[str, Any]) -> bool:
    return _currency(queue.get("debit_credit_flag")) in {"C", "CR", "CREDIT"}


def _effective_debit(transaction: "ClearingTransactionInput | dict[str, Any]") -> Decimal:
    value = transaction.clear_amount if isinstance(transaction, ClearingTransactionInput) else transaction.get("clear_amount")
    if value not in (None, ""):
        return _money(value)
    amount = transaction.amount if isinstance(transaction, ClearingTransactionInput) else transaction.get("amount", 0)
    charge = transaction.charge if isinstance(transaction, ClearingTransactionInput) else transaction.get("charge", 0)
    return _money(amount) + _money(charge)


def clearing_operating_window(at: datetime | None = None) -> dict[str, Any]:
    zone_name = settings.NEXUS_CLEARING_TIMEZONE
    current = (at or datetime.now(timezone.utc)).astimezone(ZoneInfo(zone_name))
    blocked = time(8, 0) <= current.time().replace(tzinfo=None) < time(19, 0)
    return {
        "timezone": zone_name,
        "local_time": current.isoformat(),
        "blocked": blocked,
        "batch_threshold": CLEARING_TRANCHE_SIZE,
        "window": "08:00-19:00",
        "message": (
            "Up to 20 armed records can continue during business hours. Larger sources remain available in rolling tranches."
            if blocked
            else "The unlimited command window is open; every pending source row may be armed together."
        ),
    }


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalized_text(value)
    formats = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y")
    for date_format in formats:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError as exc:
        raise ValueError(f"Invalid transaction date: {value!r}") from exc


class ClearingTransactionInput(BaseModel):
    source_row: int = Field(ge=1)
    entry_date: date
    from_account: str
    source_internal_account: str | None = Field(default=None, max_length=64)
    to_account: str = ""
    currency: str | None = Field(default=None, min_length=3, max_length=3)
    rrn: str
    stan: str
    amount: Decimal = Field(ge=0)
    charge: Decimal = Field(default=Decimal("0"), ge=0)
    clear_amount: Decimal | None = Field(default=None, ge=0)
    operation_mode: Literal[
        "QUEUE_EXACT",
        "QUEUE_PARTIAL",
        "BALANCE_ALL",
        "QUEUE_ROW_ONLY",
        "QUEUE_AMOUNT_RESET",
    ] = "QUEUE_EXACT"
    queue_row_id: str | None = Field(default=None, max_length=64)
    narration: str = Field(default="", max_length=1000)

    @field_validator("from_account", mode="before")
    @classmethod
    def validate_account(cls, value: Any) -> str:
        return _normalized_account(value)

    @field_validator("source_internal_account", mode="before")
    @classmethod
    def validate_source_internal_account(cls, value: Any) -> str | None:
        text = _normalized_text(value)
        return _normalized_account(text) if text else None

    @field_validator("to_account", mode="before")
    @classmethod
    def validate_optional_account(cls, value: Any) -> str:
        text = _normalized_text(value)
        return _normalized_account(text) if text else ""

    @field_validator("currency", mode="before")
    @classmethod
    def validate_currency(cls, value: Any) -> str | None:
        text = _normalized_text(value).upper()
        if not text:
            return None
        if not re.fullmatch(r"[A-Z0-9]{3}", text):
            raise ValueError("Currency must be a three-character currency code.")
        return text

    @field_validator("rrn", "stan", mode="before")
    @classmethod
    def validate_reference(cls, value: Any, info) -> str:
        return _normalized_reference(value, info.field_name.upper())

    @field_validator("amount", "charge", "clear_amount", mode="before")
    @classmethod
    def validate_money(cls, value: Any) -> Decimal | None:
        if value in (None, ""):
            return None
        return _money(value)

    @model_validator(mode="after")
    def validate_zero_value_placeholder(self) -> "ClearingTransactionInput":
        zero_value_modes = {"BALANCE_ALL", "QUEUE_ROW_ONLY", "QUEUE_AMOUNT_RESET"}
        if self.operation_mode not in zero_value_modes and self.amount <= 0:
            raise ValueError("Queue-linked clearing transactions must have a positive amount.")
        if (
            self.operation_mode not in zero_value_modes
            and self.clear_amount is not None
            and self.clear_amount <= 0
        ):
            raise ValueError("Queue-linked clearing transactions must have a positive clear amount.")
        return self


class ClearingBatchCreateRequest(BaseModel):
    batch_name: str = Field(min_length=3, max_length=160)
    finance_reference: str = Field(min_length=2, max_length=160)
    source_filename: str = Field(default="manual-entry.json", min_length=1, max_length=255)
    source_kind: Literal["FINANCE_IMPORT", "IDENTITY_IMPORT", "ACCOUNT_EXPLORER"] = "FINANCE_IMPORT"
    transactions: list[ClearingTransactionInput] = Field(min_length=1)


class ClearingSourceImportRequest(BaseModel):
    batch_name: str = Field(min_length=3, max_length=160)
    finance_reference: str = Field(min_length=2, max_length=160)
    filename: str = Field(min_length=1, max_length=255)
    content_base64: str = Field(min_length=1, max_length=12_000_000)
    identity_kind: Literal["ACCOUNT", "RRN"]


class ClearingSelectionRequest(BaseModel):
    selected: bool


class ClearingBulkSelectionRequest(BaseModel):
    action: Literal["ARM_WINDOW", "DISARM_ALL"] = "ARM_WINDOW"


class ClearingSubmitRequest(BaseModel):
    change_reference: str = Field(min_length=3, max_length=240)
    note: str = Field(default="", max_length=1000)


class ClearingApprovalRequest(BaseModel):
    approve: bool
    note: str = Field(default="", max_length=1000)
    idempotency_key: str | None = Field(default=None, max_length=180)

    @model_validator(mode="after")
    def validate_decision_note(self) -> "ClearingApprovalRequest":
        if len(self.note.strip()) < 3:
            raise ValueError("Record a short checker note before making this decision.")
        return self


class ClearingPolicyUpdateRequest(BaseModel):
    enabled: bool


class ClearingExecutionRequest(BaseModel):
    reason: str = Field(min_length=8, max_length=1000)
    idempotency_key: str | None = Field(default=None, max_length=180)


class ClearingRollbackRequest(BaseModel):
    reason: str = Field(min_length=12, max_length=1000)


class ClearingAccountBatchRequest(BaseModel):
    balance_row_id: str = Field(min_length=1, max_length=64)
    currency: str = Field(min_length=3, max_length=3)
    mode: Literal[
        "QUEUE_TRANSACTION",
        "ALL_UNAUTHORIZED_DEBITS",
        "QUEUE_ROW_ONLY",
        "QUEUE_AMOUNT_RESET",
    ]
    queue_row_id: str | None = Field(default=None, max_length=64)
    queue_row_ids: list[str] = Field(default_factory=list, max_length=100)
    clear_amount: Decimal | None = Field(default=None, gt=0)
    change_reference: str = Field(min_length=3, max_length=160)
    note: str = Field(default="", max_length=1000)

    @field_validator("currency", mode="before")
    @classmethod
    def validate_currency(cls, value: Any) -> str:
        text = _normalized_text(value).upper()
        if not re.fullmatch(r"[A-Z0-9]{3}", text):
            raise ValueError("Currency must be a three-character currency code.")
        return text

    @field_validator("clear_amount", mode="before")
    @classmethod
    def validate_clear_amount(cls, value: Any) -> Decimal | None:
        if value in (None, ""):
            return None
        return _money(value)

    @model_validator(mode="after")
    def normalize_queue_rows(self) -> "ClearingAccountBatchRequest":
        row_ids = [
            _normalized_text(value)
            for value in ([self.queue_row_id] if self.queue_row_id else []) + self.queue_row_ids
            if _normalized_text(value)
        ]
        self.queue_row_ids = list(dict.fromkeys(row_ids))
        self.queue_row_id = self.queue_row_ids[0] if len(self.queue_row_ids) == 1 else None
        if self.mode in {"QUEUE_TRANSACTION", "QUEUE_ROW_ONLY"} and not self.queue_row_ids:
            raise ValueError("Select at least one exact queue transaction for this custody action.")
        return self


def transaction_fingerprint(transaction: ClearingTransactionInput | dict[str, Any]) -> str:
    if isinstance(transaction, ClearingTransactionInput):
        payload = transaction.model_dump()
    else:
        payload = dict(transaction)
    canonical = {
        "entry_date": _parse_date(payload["entry_date"]).isoformat(),
        "from_account": _normalized_account(payload["from_account"]),
        "to_account": _normalized_text(payload.get("to_account")),
        "currency": _normalized_text(payload.get("currency")).upper(),
        "rrn": _normalized_reference(payload["rrn"], "RRN"),
        "stan": _normalized_reference(payload["stan"], "STAN"),
        "amount": format(_money(payload["amount"]), "f"),
        "charge": format(_money(payload.get("charge", 0)), "f"),
        "clear_amount": (
            format(_money(payload["clear_amount"]), "f")
            if payload.get("clear_amount") not in (None, "")
            else ""
        ),
        "operation_mode": _normalized_text(payload.get("operation_mode") or "QUEUE_EXACT"),
        "queue_row_id": _normalized_text(payload.get("queue_row_id")),
    }
    return hashlib.sha256(_json(canonical).encode("utf-8")).hexdigest()


HEADER_ALIASES = {
    "date": "entry_date",
    "entrydate": "entry_date",
    "transactiondate": "entry_date",
    "from": "from_account",
    "fromaccount": "from_account",
    "fromaccountnumber": "from_account",
    "debitaccount": "from_account",
    "debitaccountnumber": "from_account",
    "sourceaccount": "from_account",
    "sourceaccountnumber": "from_account",
    "to": "to_account",
    "toaccount": "to_account",
    "toaccountnumber": "to_account",
    "creditaccount": "to_account",
    "creditaccountnumber": "to_account",
    "destinationaccount": "to_account",
    "destinationaccountnumber": "to_account",
    "currency": "currency",
    "currencycode": "currency",
    "transactioncurrency": "currency",
    "ccy": "currency",
    "rrn": "rrn",
    "retrievalreferencenumber": "rrn",
    "retrievalreferenceno": "rrn",
    "referenceretrievalnumber": "rrn",
    "stan": "stan",
    "stanno": "stan",
    "stannumber": "stan",
    "transactionstan": "stan",
    "tracenumber": "stan",
    "traceno": "stan",
    "systemtracenumber": "stan",
    "systemtraceauditnumber": "stan",
    "systemtraceauditno": "stan",
    "amount": "amount",
    "transactionamount": "amount",
    "charge": "charge",
    "fee": "charge",
    "feeamount": "charge",
    "charges": "charge",
    "narration": "narration",
    "description": "narration",
}

REQUIRED_SOURCE_FIELDS = {
    "entry_date": "Date",
    "from_account": "From Account",
    "rrn": "RRN",
    "stan": "STAN",
    "amount": "Amount",
}


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", _normalized_text(value).lower())


def _rows_from_source(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        text = content.decode("utf-8-sig")
        return [dict(row) for row in csv.DictReader(io.StringIO(text))]
    if suffix in {"xlsx", "xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX import requires openpyxl. Install the Nexus production requirements.") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                return []
            return [
                {str(headers[index] or ""): row[index] if index < len(row) else None for index in range(len(headers))}
                for row in values
                if any(value not in (None, "") for value in row)
            ]
        finally:
            workbook.close()
    raise ValueError("Use a CSV or XLSX Finance source file.")


def parse_finance_source(filename: str, content: bytes) -> list[ClearingTransactionInput]:
    raw_rows = _rows_from_source(filename, content)
    if not raw_rows:
        raise ValueError("The Finance source file contains no transactions.")
    recognized_fields = {
        target
        for key in raw_rows[0]
        if (target := HEADER_ALIASES.get(_header_key(key)))
    }
    missing_fields = [
        label
        for field, label in REQUIRED_SOURCE_FIELDS.items()
        if field not in recognized_fields
    ]
    if missing_fields:
        expected = ", ".join(REQUIRED_SOURCE_FIELDS.values())
        raise ValueError(
            "Finance source is missing required column(s): "
            f"{', '.join(missing_fields)}. Required columns: {expected}. "
            "Optional columns: To Account, Currency, Fee, Narration."
        )
    transactions: list[ClearingTransactionInput] = []
    errors: list[str] = []
    for index, raw in enumerate(raw_rows, start=2):
        normalized: dict[str, Any] = {"source_row": index}
        for key, value in raw.items():
            target = HEADER_ALIASES.get(_header_key(key))
            if target:
                normalized[target] = value
        normalized.setdefault("charge", 0)
        normalized.setdefault("narration", "")
        try:
            normalized["entry_date"] = _parse_date(normalized.get("entry_date"))
            transactions.append(ClearingTransactionInput.model_validate(normalized))
        except ValidationError as exc:
            messages = []
            for error in exc.errors(include_url=False, include_input=False):
                field = str(error.get("loc", ["field"])[-1])
                label = REQUIRED_SOURCE_FIELDS.get(field, field.replace("_", " ").title())
                message = str(error.get("msg") or "is invalid")
                if error.get("type") == "missing":
                    message = "is blank"
                elif message.lower().startswith("value error, "):
                    message = message[13:]
                messages.append(f"{label} {message}")
            errors.append(f"Row {index}: {', '.join(messages)}")
        except Exception as exc:
            errors.append(f"Row {index}: {exc}")
    if errors:
        preview = "; ".join(errors[:8])
        suffix = f" and {len(errors) - 8} more" if len(errors) > 8 else ""
        raise ValueError(f"Finance source validation failed: {preview}{suffix}.")
    return transactions


def parse_identity_source(filename: str, content: bytes, identity_kind: Literal["ACCOUNT", "RRN"]) -> list[str]:
    raw_rows = _rows_from_source(filename, content)
    if not raw_rows:
        raise ValueError("The identity source file contains no records.")
    source_headers = [value for value in raw_rows[0].keys() if str(value or "").strip()]
    headers = [str(value).strip() for value in source_headers]
    accepted = {
        "ACCOUNT": {"account", "accountnumber", "debitaccount", "fromaccount"},
        "RRN": {"rrn", "retrievalreferencenumber", "retrievalreferenceno"},
    }[identity_kind]
    if len(headers) != 1 or _header_key(headers[0]) not in accepted:
        expected = "Account" if identity_kind == "ACCOUNT" else "RRN"
        raise ValueError(f"Use a one-column file with the header {expected}. Remove every other column.")

    values: list[str] = []
    errors: list[str] = []
    source_header = source_headers[0]
    for index, raw in enumerate(raw_rows, start=2):
        value = raw.get(source_header)
        try:
            normalized = (
                _normalized_account(value)
                if identity_kind == "ACCOUNT"
                else _normalized_rrn_lookup(value)
            )
            values.append(normalized)
        except Exception as exc:
            errors.append(f"Row {index}: {exc}")
    if errors:
        preview = "; ".join(errors[:8])
        suffix = f" and {len(errors) - 8} more" if len(errors) > 8 else ""
        raise ValueError(f"Identity source validation failed: {preview}{suffix}.")
    duplicates = [value for value, count in Counter(values).items() if count > 1]
    if duplicates:
        raise ValueError(f"The source contains duplicate {identity_kind} values: {', '.join(duplicates[:5])}.")
    return values


class ClearingStorageError(RuntimeError):
    pass


class ClearingRepository:
    """Postgres custody ledger for the complete clearing lifecycle."""

    def __init__(self) -> None:
        self._dsn = settings.nexus_database_dsn
        self._local_locks: dict[str, threading.Lock] = {}

    @property
    def database_dsn(self) -> str | None:
        return self._dsn

    @contextmanager
    def _connection(self):
        if not self._dsn:
            raise ClearingStorageError("Unauthorized clearing requires the shared SentinelOps database.")
        try:
            with psycopg.connect(self._dsn, row_factory=dict_row) as connection:
                yield connection
        except (psycopg.errors.UndefinedTable, psycopg.errors.UndefinedColumn) as exc:
            raise ClearingStorageError(
                f"Unauthorized clearing storage is not initialized. Apply {CLEARING_MIGRATION} and restart Nexus."
            ) from exc
        except psycopg.errors.InsufficientPrivilege as exc:
            raise ClearingStorageError(
                "The Nexus database role cannot access unauthorized clearing storage. "
                "Grant SELECT, INSERT, UPDATE, and the required sequence privileges on nexus_clearing_* tables."
            ) from exc

    def create_batch(
        self,
        request: ClearingBatchCreateRequest,
        *,
        actor: str,
        actor_role: str,
        source_sha256: str | None = None,
    ) -> dict[str, Any]:
        if len(request.transactions) > settings.NEXUS_CLEARING_MAX_BATCH_SIZE:
            raise ValueError(
                f"A clearing batch cannot exceed {settings.NEXUS_CLEARING_MAX_BATCH_SIZE} transactions."
            )
        fingerprints = [transaction_fingerprint(item) for item in request.transactions]
        duplicates = [fingerprint for fingerprint, count in Counter(fingerprints).items() if count > 1]
        if duplicates:
            raise ValueError("The Finance source contains duplicate authorized transactions.")
        queue_keys = [
            (
                item.entry_date,
                item.from_account,
                item.rrn,
                item.stan,
                item.amount,
                item.charge,
            )
            for item in request.transactions
        ]
        if any(count > 1 for count in Counter(queue_keys).values()):
            raise ValueError(
                "The Finance source assigns the same Oracle queue identity to more than one transaction."
            )
        normalized_source = [
            {**item.model_dump(), "fingerprint": fingerprint}
            for item, fingerprint in zip(request.transactions, fingerprints)
        ]
        source_hash = source_sha256 or hashlib.sha256(_json(normalized_source).encode("utf-8")).hexdigest()
        batch_id = f"nexus-clear-batch-{uuid4()}"
        total_amount = sum((item.amount for item in request.transactions), Decimal("0"))
        total_charge = sum((item.charge for item in request.transactions), Decimal("0"))
        total_debit = sum((_effective_debit(item) for item in request.transactions), Decimal("0"))
        debit_accounts = {item.from_account for item in request.transactions}
        operating_window = clearing_operating_window()
        armed_count = (
            min(len(request.transactions), CLEARING_TRANCHE_SIZE)
            if operating_window["blocked"]
            else len(request.transactions)
        )
        try:
            with self._connection() as connection:
                with connection.cursor() as cursor:
                    # A deleted batch with no approval or execution is only retired
                    # intake/read evidence. Release its transaction identities so a
                    # corrected source can be imported while preserving the old batch,
                    # reconciliation snapshots, and audit chronology.
                    cursor.execute(
                        """
                        DELETE FROM nexus_clearing_transaction transaction
                        USING nexus_clearing_batch batch
                        WHERE transaction.batch_id = batch.batch_id
                          AND transaction.fingerprint = ANY(%s)
                          AND batch.deleted_at IS NOT NULL
                          AND NOT EXISTS (
                              SELECT 1 FROM nexus_clearing_approval approval
                              WHERE approval.batch_id = batch.batch_id
                          )
                          AND NOT EXISTS (
                              SELECT 1 FROM nexus_clearing_execution_run execution
                              WHERE execution.batch_id = batch.batch_id
                          )
                        """,
                        (fingerprints,),
                    )
                    cursor.execute(
                        """
                        INSERT INTO nexus_clearing_batch (
                            batch_id, batch_name, finance_reference, source_filename, source_sha256,
                            transaction_count, debit_account_count, credit_account_count,
                            total_amount, total_charge, total_debit, selected_count, created_by, source_kind
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        """,
                        (
                            batch_id,
                            request.batch_name.strip(),
                            request.finance_reference.strip(),
                            request.source_filename,
                            source_hash,
                            len(request.transactions),
                            len(debit_accounts),
                            0,
                            total_amount,
                            total_charge,
                            total_debit,
                            armed_count,
                            actor,
                            request.source_kind,
                        ),
                    )
                    for index, (item, fingerprint) in enumerate(zip(request.transactions, fingerprints)):
                        armed = index < armed_count
                        cursor.execute(
                            """
                            INSERT INTO nexus_clearing_transaction (
                                fingerprint, batch_id, source_row, entry_date, from_account,
                                source_internal_account, to_account,
                                currency, rrn, stan, amount, charge, clear_amount, operation_mode,
                                queue_row_id, narration, selected, reconciliation_state, custody_state
                            )
                            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'PENDING')
                            """,
                            (
                                fingerprint,
                                batch_id,
                                item.source_row,
                                item.entry_date,
                                item.from_account,
                                item.source_internal_account,
                                item.to_account,
                                item.currency,
                                item.rrn,
                                item.stan,
                                item.amount,
                                item.charge,
                                item.clear_amount,
                                item.operation_mode,
                                item.queue_row_id,
                                item.narration,
                                armed,
                                "NOT_RECONCILED" if armed else "EXCLUDED",
                            ),
                        )
                    self._audit_cursor(
                        cursor,
                        event_type="batch_imported",
                        actor=actor,
                        actor_role=actor_role,
                        batch_id=batch_id,
                        details={
                            "source_filename": request.source_filename,
                            "source_sha256": source_hash,
                            "finance_reference": request.finance_reference,
                            "transaction_count": len(request.transactions),
                            "armed_count": armed_count,
                            "remaining_count": len(request.transactions) - armed_count,
                            "total_debit": total_debit,
                            "source_kind": request.source_kind,
                        },
                    )
                connection.commit()
        except psycopg.errors.UniqueViolation as exc:
            raise ValueError(
                "This source is already active or entered approval/execution custody. Delete a non-executed intake before re-importing it, "
                "or use a new Finance reference for a true revision."
            ) from exc
        except psycopg.errors.CheckViolation as exc:
            if exc.diag.constraint_name == "nexus_clearing_batch_source_kind_check":
                raise ClearingStorageError(
                    "Funds Custody identity intake is newer than the database contract. "
                    "Apply 2026_12_repair_nexus_clearing_source_contract.sql and restart Nexus."
                ) from exc
            if exc.diag.constraint_name in {
                "nexus_clearing_transaction_operation_mode_check",
                "nexus_clearing_transaction_clear_amount_check",
                "nexus_clearing_transaction_amount_check",
            }:
                raise ClearingStorageError(
                    "Funds Custody queue-only actions are newer than the database contract. "
                    "Apply 2026_18_add_nexus_clearing_queue_custody.sql and restart Nexus."
                ) from exc
            raise
        return self.get_batch(batch_id)

    def list_batches(self, limit: int = 50) -> list[dict[str, Any]]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                rows = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_batch
                    WHERE deleted_at IS NULL
                    ORDER BY updated_at DESC
                    LIMIT %s
                    """,
                    (max(1, min(limit, 200)),),
                ).fetchall()
        return [_plain(dict(row)) for row in rows]

    def delete_batch(self, batch_id: str, *, actor: str, actor_role: str) -> None:
        deletable = {"IMPORTED", "RECONCILED", "HAS_EXCEPTIONS", "READY_FOR_APPROVAL", "FAILED", "BLOCKED"}
        account_explorer_terminal = {"COMPLETED", "ROLLED_BACK"}
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    """
                    SELECT status, source_kind,
                           EXISTS (SELECT 1 FROM nexus_clearing_approval WHERE batch_id = %s) AS submitted,
                           EXISTS (SELECT 1 FROM nexus_clearing_execution_run WHERE batch_id = %s) AS executed
                    FROM nexus_clearing_batch
                    WHERE batch_id = %s AND deleted_at IS NULL
                    FOR UPDATE
                    """,
                    (batch_id, batch_id, batch_id),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                is_terminal_account_batch = (
                    batch["source_kind"] == "ACCOUNT_EXPLORER"
                    and batch["status"] in account_explorer_terminal
                )
                if batch["status"] not in deletable and not is_terminal_account_batch:
                    raise ValueError("A batch in approval, execution, or completed custody cannot be deleted.")
                self._audit_cursor(
                    cursor,
                    event_type="batch_archived" if is_terminal_account_batch else "batch_deleted",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    details={
                        "previous_status": batch["status"],
                        "source_kind": batch["source_kind"],
                        "custody_preserved": bool(batch["submitted"] or batch["executed"]),
                    },
                )
                cursor.execute(
                    "UPDATE nexus_clearing_batch SET deleted_at = now(), deleted_by = %s, updated_at = now() WHERE batch_id = %s",
                    (actor, batch_id),
                )
                if not batch["submitted"] and not batch["executed"]:
                    cursor.execute(
                        "DELETE FROM nexus_clearing_transaction WHERE batch_id = %s",
                        (batch_id,),
                    )
            connection.commit()

    def overview(self) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                row = cursor.execute(
                    """
                    SELECT
                        COUNT(*) FILTER (WHERE status NOT IN ('COMPLETED', 'ROLLED_BACK')) AS open_batches,
                        COUNT(*) FILTER (WHERE status = 'PENDING_APPROVAL') AS awaiting_approval,
                        COUNT(*) FILTER (WHERE status IN ('APPROVED', 'EXECUTION_READY')) AS ready_to_execute,
                        MAX(updated_at) AS last_activity_at
                    FROM nexus_clearing_batch
                    WHERE deleted_at IS NULL
                    """
                ).fetchone()
                latest = cursor.execute(
                    """
                    SELECT batch_id, batch_name, finance_reference, status, transaction_count, selected_count,
                           total_debit, updated_at
                    FROM nexus_clearing_batch
                    WHERE deleted_at IS NULL
                    ORDER BY updated_at DESC
                    LIMIT 1
                    """
                ).fetchone()
        return {
            **_plain(dict(row or {})),
            "latest_batch": _plain(dict(latest)) if latest else None,
            "writes_enabled": settings.NEXUS_CLEARING_WRITES_ENABLED,
            "production_writes_enabled": settings.NEXUS_CLEARING_PRODUCTION_WRITES_ENABLED,
            "environment": settings.ENVIRONMENT,
        }

    def get_specific_record_policy(self) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                row = cursor.execute(
                    """
                    SELECT policy_key, enabled, updated_by, updated_at
                    FROM nexus_clearing_policy
                    WHERE policy_key = 'specific-record-clearing'
                    """
                ).fetchone()
        if not row:
            raise ClearingStorageError(
                "The specific-record clearing policy is not initialized. "
                f"Apply {CLEARING_MIGRATION} and restart Nexus."
            )
        return _plain(dict(row))

    def set_specific_record_policy(
        self,
        enabled: bool,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                previous = cursor.execute(
                    """
                    SELECT enabled
                    FROM nexus_clearing_policy
                    WHERE policy_key = 'specific-record-clearing'
                    FOR UPDATE
                    """
                ).fetchone()
                if not previous:
                    raise ClearingStorageError(
                        "The specific-record clearing policy is not initialized. "
                        f"Apply {CLEARING_MIGRATION} and restart Nexus."
                    )
                cursor.execute(
                    """
                    UPDATE nexus_clearing_policy
                    SET enabled = %s, updated_by = %s, updated_at = now()
                    WHERE policy_key = 'specific-record-clearing'
                    """,
                    (enabled, actor),
                )
                self._audit_cursor(
                    cursor,
                    event_type="specific_record_policy_changed",
                    actor=actor,
                    actor_role=actor_role,
                    details={"previous_enabled": previous["enabled"], "enabled": enabled},
                )
            connection.commit()
        return self.get_specific_record_policy()

    def get_batch(self, batch_id: str) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    "SELECT * FROM nexus_clearing_batch WHERE batch_id = %s AND deleted_at IS NULL",
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                transactions = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_transaction
                    WHERE batch_id = %s
                    ORDER BY source_row
                    """,
                    (batch_id,),
                ).fetchall()
                approvals = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_approval
                    WHERE batch_id = %s
                    ORDER BY submitted_at DESC
                    """,
                    (batch_id,),
                ).fetchall()
                reconciliation = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_reconciliation_run
                    WHERE batch_id = %s
                    ORDER BY started_at DESC
                    LIMIT 1
                    """,
                    (batch_id,),
                ).fetchone()
        plain_transactions = [_plain(dict(row)) for row in transactions]
        committed_count = sum(1 for row in plain_transactions if row.get("custody_state") == "COMMITTED")
        remaining_count = sum(1 for row in plain_transactions if row.get("custody_state", "PENDING") == "PENDING")
        armed_count = sum(
            1
            for row in plain_transactions
            if row.get("custody_state", "PENDING") == "PENDING" and row.get("selected")
        )
        return {
            **_plain(dict(batch)),
            "transactions": plain_transactions,
            "approvals": [_plain(dict(row)) for row in approvals],
            "latest_reconciliation": _plain(dict(reconciliation)) if reconciliation else None,
            "command_scope": {
                "limit": CLEARING_TRANCHE_SIZE,
                "armed": armed_count,
                "committed": committed_count,
                "remaining": remaining_count,
                "complete": remaining_count == 0,
            },
        }

    def selected_transactions(self, batch_id: str) -> list[dict[str, Any]]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                rows = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_transaction
                    WHERE batch_id = %s AND selected = TRUE AND custody_state = 'PENDING'
                    ORDER BY source_row
                    """,
                    (batch_id,),
                ).fetchall()
        return [dict(row) for row in rows]

    @staticmethod
    def _reset_after_selection_cursor(cursor, batch_id: str) -> int:
        cursor.execute(
            """
            UPDATE nexus_clearing_transaction
            SET reconciliation_state = CASE
                    WHEN custody_state = 'COMMITTED' THEN 'NO_LONGER_OUTSTANDING'
                    WHEN selected THEN 'STALE'
                    ELSE 'EXCLUDED'
                END,
                reconciliation_payload = '{}'::jsonb,
                updated_at = now()
            WHERE batch_id = %s
            """,
            (batch_id,),
        )
        selected_count = cursor.execute(
            "SELECT COUNT(*) AS count FROM nexus_clearing_transaction WHERE batch_id = %s AND selected AND custody_state = 'PENDING'",
            (batch_id,),
        ).fetchone()["count"]
        cursor.execute(
            """
            UPDATE nexus_clearing_batch
            SET status = 'IMPORTED',
                selected_count = %s,
                reconciliation_summary = '{}'::jsonb,
                approved_payload = NULL,
                payload_hash = NULL,
                submitted_by = NULL,
                submitted_at = NULL,
                approved_by = NULL,
                approved_at = NULL,
                updated_at = now()
            WHERE batch_id = %s
            """,
            (selected_count, batch_id),
        )
        return int(selected_count)

    def set_selection(
        self,
        batch_id: str,
        fingerprint: str,
        selected: bool,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        window = clearing_operating_window()
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    "SELECT status FROM nexus_clearing_batch WHERE batch_id = %s AND deleted_at IS NULL FOR UPDATE",
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                if batch["status"] in {"EXECUTING", "ROLLED_BACK", "COMMIT_UNCERTAIN"}:
                    raise ValueError("Selection is locked after execution begins.")
                transaction = cursor.execute(
                    """
                    SELECT selected, custody_state
                    FROM nexus_clearing_transaction
                    WHERE batch_id = %s AND fingerprint = %s
                    FOR UPDATE
                    """,
                    (batch_id, fingerprint),
                ).fetchone()
                if not transaction:
                    raise LookupError("Finance transaction not found in this batch.")
                if transaction["custody_state"] == "COMMITTED":
                    raise ValueError("A committed source row is immutable and cannot re-enter selection.")
                if selected and not transaction["selected"]:
                    armed = cursor.execute(
                        """
                        SELECT COUNT(*) AS count
                        FROM nexus_clearing_transaction
                        WHERE batch_id = %s AND selected AND custody_state = 'PENDING'
                        """,
                        (batch_id,),
                    ).fetchone()["count"]
                    if window["blocked"] and armed >= CLEARING_TRANCHE_SIZE:
                        raise ValueError(
                            f"The active command tranche is full at {CLEARING_TRANCHE_SIZE} records. "
                            "Complete or exclude a selected row before arming another."
                        )
                updated = cursor.execute(
                    """
                    UPDATE nexus_clearing_transaction
                    SET selected = %s, updated_at = now()
                    WHERE batch_id = %s AND fingerprint = %s AND custody_state = 'PENDING'
                    RETURNING fingerprint
                    """,
                    (selected, batch_id, fingerprint),
                ).fetchone()
                if not updated:
                    raise LookupError("Finance transaction not found in this batch.")
                selected_count = self._reset_after_selection_cursor(cursor, batch_id)
                self._notify_cursor(
                    cursor,
                    event_type="selection_changed",
                    actor=actor,
                    batch_id=batch_id,
                    fingerprint=fingerprint,
                    details={"selected": selected, "selected_count": selected_count},
                )
            connection.commit()
        return {
            "batch_id": batch_id,
            "fingerprint": fingerprint,
            "selected": selected,
            "selected_count": selected_count,
            "window_limited": window["blocked"],
            "selection_limit": CLEARING_TRANCHE_SIZE if window["blocked"] else None,
        }

    def set_bulk_selection(
        self,
        batch_id: str,
        action: str,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        window = clearing_operating_window()
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    "SELECT status FROM nexus_clearing_batch WHERE batch_id = %s AND deleted_at IS NULL FOR UPDATE",
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                if batch["status"] in {"EXECUTING", "COMPLETED", "ROLLED_BACK", "COMMIT_UNCERTAIN"}:
                    raise ValueError("Selection is locked after execution begins.")
                pending = cursor.execute(
                    """
                    SELECT fingerprint
                    FROM nexus_clearing_transaction
                    WHERE batch_id = %s AND custody_state = 'PENDING'
                    ORDER BY source_row, fingerprint
                    """,
                    (batch_id,),
                ).fetchall()
                pending_fingerprints = [str(row["fingerprint"]) for row in pending]
                if action == "DISARM_ALL":
                    selected_fingerprints: list[str] = []
                elif action == "ARM_WINDOW":
                    selected_fingerprints = (
                        pending_fingerprints[:CLEARING_TRANCHE_SIZE]
                        if window["blocked"]
                        else pending_fingerprints
                    )
                else:
                    raise ValueError("Unsupported selection action.")
                cursor.execute(
                    """
                    UPDATE nexus_clearing_transaction
                    SET selected = fingerprint = ANY(%s), updated_at = now()
                    WHERE batch_id = %s AND custody_state = 'PENDING'
                    """,
                    (selected_fingerprints, batch_id),
                )
                selected_count = self._reset_after_selection_cursor(cursor, batch_id)
                self._notify_cursor(
                    cursor,
                    event_type="selection_changed",
                    actor=actor,
                    batch_id=batch_id,
                    details={
                        "action": action,
                        "selected_count": selected_count,
                        "window_limited": window["blocked"],
                    },
                )
            connection.commit()
        return {
            "batch_id": batch_id,
            "selected_fingerprints": selected_fingerprints,
            "selected_count": selected_count,
            "window_limited": window["blocked"],
            "selection_limit": CLEARING_TRANCHE_SIZE if window["blocked"] else None,
        }

    def begin_reconciliation(self, batch_id: str, actor: str) -> str:
        reconciliation_id = f"nexus-clear-reconcile-{uuid4()}"
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    "SELECT status FROM nexus_clearing_batch WHERE batch_id = %s FOR UPDATE",
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                if batch["status"] in {"EXECUTING", "COMPLETED", "ROLLED_BACK", "COMMIT_UNCERTAIN"}:
                    raise ValueError("This batch can no longer be reconciled.")
                selected_count = cursor.execute(
                    "SELECT COUNT(*) AS count FROM nexus_clearing_transaction WHERE batch_id = %s AND selected AND custody_state = 'PENDING'",
                    (batch_id,),
                ).fetchone()["count"]
                if not selected_count:
                    raise ValueError("Select at least one Finance transaction before reconciliation.")
                cursor.execute(
                    """
                    INSERT INTO nexus_clearing_reconciliation_run (
                        reconciliation_id, batch_id, status, requested_by, selected_count
                    )
                    VALUES (%s, %s, 'RUNNING', %s, %s)
                    """,
                    (reconciliation_id, batch_id, actor, selected_count),
                )
                cursor.execute(
                    "UPDATE nexus_clearing_batch SET status = 'RECONCILING', updated_at = now() WHERE batch_id = %s",
                    (batch_id,),
                )
            connection.commit()
        return reconciliation_id

    def finish_reconciliation(
        self,
        batch_id: str,
        reconciliation_id: str,
        *,
        results: dict[str, dict[str, Any]],
        snapshots: list[dict[str, Any]],
        summary: dict[str, Any],
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        blocking = int(summary.get("blocked", 0)) + int(summary.get("special_review", 0))
        ready = sum(int(summary.get(key, 0)) for key in ("ready", "ready_with_residual", "safe_extra_queue"))
        status = "HAS_EXCEPTIONS" if blocking else ("READY_FOR_APPROVAL" if ready else "RECONCILED")
        with self._connection() as connection:
            with connection.cursor() as cursor:
                for fingerprint, result in results.items():
                    selected = result["state"] != "NO_LONGER_OUTSTANDING"
                    resolved_clear_amount = result.get("clear_amount")
                    resolved_currency = result.get("currency")
                    cursor.execute(
                        """
                        UPDATE nexus_clearing_transaction
                        SET selected = %s,
                            reconciliation_state = %s,
                            reconciliation_payload = %s::jsonb,
                            amount = CASE
                                WHEN operation_mode = 'BALANCE_ALL' AND amount = 0 AND %s::numeric > 0
                                THEN %s::numeric
                                ELSE amount
                            END,
                            clear_amount = CASE
                                WHEN operation_mode = 'BALANCE_ALL' AND %s::numeric >= 0
                                THEN %s::numeric
                                ELSE clear_amount
                            END,
                            currency = COALESCE(%s, currency),
                            updated_at = now()
                        WHERE batch_id = %s AND fingerprint = %s
                          AND custody_state = 'PENDING'
                        """,
                        (
                            selected,
                            result["state"],
                            _json(result),
                            resolved_clear_amount or 0,
                            resolved_clear_amount or 0,
                            resolved_clear_amount or 0,
                            resolved_clear_amount or 0,
                            resolved_currency,
                            batch_id,
                            fingerprint,
                        ),
                    )
                for snapshot in snapshots:
                    cursor.execute(
                        """
                        INSERT INTO nexus_clearing_account_snapshot (
                            snapshot_id, reconciliation_id, batch_id, external_account, internal_account,
                            account_role, mapping_count, balance_rows, queue_rows, resolution
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s::jsonb, %s::jsonb)
                        """,
                        (
                            f"nexus-clear-snapshot-{uuid4()}",
                            reconciliation_id,
                            batch_id,
                            snapshot["external_account"],
                            snapshot.get("internal_account"),
                            snapshot["account_role"],
                            snapshot["mapping_count"],
                            _json(snapshot.get("balance_rows", [])),
                            _json(snapshot.get("queue_rows", [])),
                            _json(snapshot.get("resolution", {})),
                        ),
                    )
                selected_count = cursor.execute(
                    "SELECT COUNT(*) AS count FROM nexus_clearing_transaction WHERE batch_id = %s AND selected AND custody_state = 'PENDING'",
                    (batch_id,),
                ).fetchone()["count"]
                cursor.execute(
                    """
                    UPDATE nexus_clearing_reconciliation_run
                    SET status = %s, completed_at = now(), summary = %s::jsonb
                    WHERE reconciliation_id = %s
                    """,
                    ("HAS_EXCEPTIONS" if blocking else "COMPLETED", _json(summary), reconciliation_id),
                )
                cursor.execute(
                    """
                    UPDATE nexus_clearing_batch
                    SET status = %s,
                        selected_count = %s,
                        reconciliation_summary = %s::jsonb,
                        approved_payload = NULL,
                        payload_hash = NULL,
                        updated_at = now()
                    WHERE batch_id = %s
                    """,
                    (status, selected_count, _json(summary), batch_id),
                )
                self._audit_cursor(
                    cursor,
                    event_type="batch_reconciled",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    reconciliation_id=reconciliation_id,
                    details=summary,
                )
            connection.commit()
        return self.get_batch(batch_id)

    def fail_reconciliation(self, batch_id: str, reconciliation_id: str, error: str) -> None:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    UPDATE nexus_clearing_reconciliation_run
                    SET status = 'FAILED', completed_at = now(), error_message = %s
                    WHERE reconciliation_id = %s
                    """,
                    (error[:2000], reconciliation_id),
                )
                cursor.execute(
                    "UPDATE nexus_clearing_batch SET status = 'FAILED', updated_at = now() WHERE batch_id = %s",
                    (batch_id,),
                )
            connection.commit()

    def save_submission(
        self,
        batch_id: str,
        *,
        payload: dict[str, Any],
        payload_hash: str,
        actor: str,
        actor_role: str,
        change_reference: str,
        note: str,
    ) -> dict[str, Any]:
        approval_id = f"nexus-clear-approval-{uuid4()}"
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    "SELECT status FROM nexus_clearing_batch WHERE batch_id = %s FOR UPDATE",
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                cursor.execute(
                    """
                    UPDATE nexus_clearing_batch
                    SET status = 'PENDING_APPROVAL',
                        approved_payload = %s::jsonb,
                        payload_hash = %s,
                        submitted_by = %s,
                        submitted_at = now(),
                        approved_by = NULL,
                        approved_at = NULL,
                        updated_at = now()
                    WHERE batch_id = %s
                    """,
                    (_json(payload), payload_hash, actor, batch_id),
                )
                cursor.execute(
                    """
                    INSERT INTO nexus_clearing_approval (
                        approval_id, batch_id, payload_hash, status, submitted_by, review_note
                    )
                    VALUES (%s, %s, %s, 'PENDING', %s, %s)
                    """,
                    (approval_id, batch_id, payload_hash, actor, note),
                )
                self._audit_cursor(
                    cursor,
                    event_type="batch_submitted",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    reconciliation_id=payload.get("reconciliation_id"),
                    details={
                        "payload_hash": payload_hash,
                        "change_reference": change_reference,
                        "note": note,
                        "transaction_count": len(payload["transactions"]),
                    },
                )
            connection.commit()
        return self.get_batch(batch_id)

    def decide_approval(
        self,
        batch_id: str,
        *,
        approve: bool,
        actor: str,
        actor_role: str,
        note: str,
        allow_self_approval: bool = False,
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                batch = cursor.execute(
                    """
                    SELECT status, submitted_by, payload_hash, approved_payload
                    FROM nexus_clearing_batch
                    WHERE batch_id = %s
                    FOR UPDATE
                    """,
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                if batch["status"] != "PENDING_APPROVAL":
                    raise ValueError("This batch is not awaiting approval.")
                if str(batch["submitted_by"]).lower() == actor.lower() and not allow_self_approval:
                    raise ValueError("Maker-checker separation prevents the submitter from approving this batch.")
                pending = cursor.execute(
                    """
                    SELECT approval_id
                    FROM nexus_clearing_approval
                    WHERE batch_id = %s AND status = 'PENDING'
                    ORDER BY submitted_at DESC
                    LIMIT 1
                    FOR UPDATE
                    """,
                    (batch_id,),
                ).fetchone()
                if not pending:
                    raise ValueError("The pending approval record is missing.")
                decision = "APPROVED" if approve else "REJECTED"
                cursor.execute(
                    """
                    UPDATE nexus_clearing_approval
                    SET status = %s, reviewed_by = %s, reviewed_at = now(), review_note = %s
                    WHERE approval_id = %s
                    """,
                    (decision, actor, note, pending["approval_id"]),
                )
                cursor.execute(
                    """
                    UPDATE nexus_clearing_batch
                    SET status = %s,
                        approved_by = CASE WHEN %s THEN %s ELSE NULL END,
                        approved_at = CASE WHEN %s THEN now() ELSE NULL END,
                        updated_at = now()
                    WHERE batch_id = %s
                    """,
                    ("EXECUTION_READY" if approve else "READY_FOR_APPROVAL", approve, actor, approve, batch_id),
                )
                self._audit_cursor(
                    cursor,
                    event_type="batch_approved" if approve else "batch_rejected",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    reconciliation_id=(batch.get("approved_payload") or {}).get("reconciliation_id"),
                    details={"payload_hash": batch["payload_hash"], "note": note},
                )
            connection.commit()
        return self.get_batch(batch_id)

    def approve_and_start_execution(
        self,
        batch_id: str,
        *,
        actor: str,
        actor_role: str,
        note: str,
        idempotency_key: str,
        allow_self_approval: bool = False,
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                existing = cursor.execute(
                    "SELECT execution_id FROM nexus_clearing_execution_run WHERE idempotency_key = %s",
                    (idempotency_key,),
                ).fetchone()
                if existing:
                    return self.get_execution(existing["execution_id"])
                batch = cursor.execute(
                    """
                    SELECT status, submitted_by, payload_hash, approved_payload
                    FROM nexus_clearing_batch
                    WHERE batch_id = %s
                    FOR UPDATE
                    """,
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                if batch["status"] != "PENDING_APPROVAL":
                    raise ValueError("This batch is not awaiting approval.")
                if str(batch["submitted_by"]).lower() == actor.lower() and not allow_self_approval:
                    raise ValueError("Maker-checker separation prevents the submitter from approving this batch.")
                pending = cursor.execute(
                    """
                    SELECT approval_id
                    FROM nexus_clearing_approval
                    WHERE batch_id = %s AND status = 'PENDING'
                    ORDER BY submitted_at DESC
                    LIMIT 1
                    FOR UPDATE
                    """,
                    (batch_id,),
                ).fetchone()
                if not pending:
                    raise ValueError("The pending approval record is missing.")
                payload = batch.get("approved_payload") or {}
                if not payload or payload.get("payload_hash") != batch.get("payload_hash"):
                    raise ValueError("The sealed approval payload is missing or no longer matches its hash.")

                execution_id = f"nexus-clear-execution-{uuid4()}"
                cursor.execute(
                    """
                    UPDATE nexus_clearing_approval
                    SET status = 'APPROVED', reviewed_by = %s, reviewed_at = now(), review_note = %s
                    WHERE approval_id = %s
                    """,
                    (actor, note, pending["approval_id"]),
                )
                cursor.execute(
                    """
                    INSERT INTO nexus_clearing_execution_run (
                        execution_id, batch_id, idempotency_key, payload_hash, status,
                        requested_by, reason, approved_payload
                    )
                    VALUES (%s, %s, %s, %s, 'CREATED', %s, %s, %s::jsonb)
                    """,
                    (
                        execution_id,
                        batch_id,
                        idempotency_key,
                        batch["payload_hash"],
                        actor,
                        note,
                        _json(payload),
                    ),
                )
                for transaction in payload.get("transactions", []):
                    cursor.execute(
                        """
                        INSERT INTO nexus_clearing_execution_item (
                            execution_item_id, execution_id, fingerprint, status
                        )
                        VALUES (%s, %s, %s, 'PENDING')
                        """,
                        (f"nexus-clear-item-{uuid4()}", execution_id, transaction["fingerprint"]),
                    )
                cursor.execute(
                    """
                    UPDATE nexus_clearing_batch
                    SET status = 'EXECUTING', approved_by = %s, approved_at = now(), updated_at = now()
                    WHERE batch_id = %s
                    """,
                    (actor, batch_id),
                )
                self._audit_cursor(
                    cursor,
                    event_type="batch_approved",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    reconciliation_id=payload.get("reconciliation_id"),
                    execution_id=execution_id,
                    details={"payload_hash": batch["payload_hash"], "note": note, "execution_started": True},
                )
                self._audit_cursor(
                    cursor,
                    event_type="execution_created",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    reconciliation_id=payload.get("reconciliation_id"),
                    execution_id=execution_id,
                    details={"idempotency_key": idempotency_key, "payload_hash": batch["payload_hash"]},
                )
            connection.commit()
        return self.get_execution(execution_id)

    def start_execution(
        self,
        batch_id: str,
        *,
        actor: str,
        actor_role: str,
        reason: str,
        idempotency_key: str,
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                existing = cursor.execute(
                    "SELECT * FROM nexus_clearing_execution_run WHERE idempotency_key = %s",
                    (idempotency_key,),
                ).fetchone()
                if existing:
                    return _plain(dict(existing))
                batch = cursor.execute(
                    """
                    SELECT status, payload_hash, approved_payload
                    FROM nexus_clearing_batch
                    WHERE batch_id = %s
                    FOR UPDATE
                    """,
                    (batch_id,),
                ).fetchone()
                if not batch:
                    raise LookupError("Clearing batch not found.")
                if batch["status"] != "EXECUTION_READY":
                    raise ValueError("The batch must be approved and execution-ready.")
                execution_id = f"nexus-clear-execution-{uuid4()}"
                cursor.execute(
                    """
                    INSERT INTO nexus_clearing_execution_run (
                        execution_id, batch_id, idempotency_key, payload_hash, status,
                        requested_by, reason, approved_payload
                    )
                    VALUES (%s, %s, %s, %s, 'CREATED', %s, %s, %s::jsonb)
                    """,
                    (
                        execution_id,
                        batch_id,
                        idempotency_key,
                        batch["payload_hash"],
                        actor,
                        reason,
                        _json(batch["approved_payload"]),
                    ),
                )
                for transaction in batch["approved_payload"].get("transactions", []):
                    cursor.execute(
                        """
                        INSERT INTO nexus_clearing_execution_item (
                            execution_item_id, execution_id, fingerprint, status
                        )
                        VALUES (%s, %s, %s, 'PENDING')
                        """,
                        (f"nexus-clear-item-{uuid4()}", execution_id, transaction["fingerprint"]),
                    )
                cursor.execute(
                    "UPDATE nexus_clearing_batch SET status = 'EXECUTING', updated_at = now() WHERE batch_id = %s",
                    (batch_id,),
                )
                self._audit_cursor(
                    cursor,
                    event_type="execution_created",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=batch_id,
                    reconciliation_id=(batch.get("approved_payload") or {}).get("reconciliation_id"),
                    execution_id=execution_id,
                    details={"idempotency_key": idempotency_key, "payload_hash": batch["payload_hash"]},
                )
            connection.commit()
        return self.get_execution(execution_id)

    def update_execution_status(self, execution_id: str, status: str) -> None:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    UPDATE nexus_clearing_execution_run
                    SET status = %s,
                        started_at = COALESCE(started_at, now())
                    WHERE execution_id = %s
                    """,
                    (status, execution_id),
                )
            connection.commit()

    def finish_execution(
        self,
        execution_id: str,
        *,
        status: str,
        evidence: dict[str, Any],
        error: str | None,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                execution = cursor.execute(
                    "SELECT batch_id, approved_payload FROM nexus_clearing_execution_run WHERE execution_id = %s FOR UPDATE",
                    (execution_id,),
                ).fetchone()
                if not execution:
                    raise LookupError("Clearing execution not found.")
                cursor.execute(
                    """
                    UPDATE nexus_clearing_execution_run
                    SET status = %s,
                        oracle_evidence = %s::jsonb,
                        error_message = %s,
                        completed_at = now()
                    WHERE execution_id = %s
                    """,
                    (status, _json(evidence), error, execution_id),
                )
                item_status = "COMMITTED" if status == "COMMITTED" else ("BLOCKED" if status == "BLOCKED" else "FAILED")
                item_evidence = _execution_item_evidence(
                    execution.get("approved_payload") or {},
                    evidence,
                )
                for fingerprint, scoped_evidence in item_evidence.items():
                    cursor.execute(
                        """
                        UPDATE nexus_clearing_execution_item
                        SET status = %s,
                            before_evidence = %s::jsonb,
                            after_evidence = %s::jsonb,
                            error_message = %s,
                            updated_at = now()
                        WHERE execution_id = %s AND fingerprint = %s
                        """,
                        (
                            item_status,
                            _json(scoped_evidence["before"]),
                            _json(scoped_evidence["after"]),
                            error,
                            execution_id,
                            fingerprint,
                        ),
                    )
                cursor.execute(
                    """
                    UPDATE nexus_clearing_execution_item
                    SET status = %s,
                        error_message = %s,
                        updated_at = now()
                    WHERE execution_id = %s AND status = 'PENDING'
                    """,
                    (item_status, error, execution_id),
                )
                remaining_count = 0
                committed_count = 0
                next_armed_count = 0
                if status == "COMMITTED":
                    cursor.execute(
                        """
                        UPDATE nexus_clearing_transaction transaction
                        SET custody_state = 'COMMITTED',
                            committed_execution_id = %s,
                            committed_at = now(),
                            selected = FALSE,
                            reconciliation_state = 'NO_LONGER_OUTSTANDING',
                            updated_at = now()
                        FROM nexus_clearing_execution_item item
                        WHERE item.execution_id = %s
                          AND item.fingerprint = transaction.fingerprint
                          AND transaction.batch_id = %s
                        """,
                        (execution_id, execution_id, execution["batch_id"]),
                    )
                    cursor.execute(
                        """
                        WITH pending AS (
                            SELECT fingerprint,
                                   ROW_NUMBER() OVER (ORDER BY source_row, fingerprint) AS command_rank
                            FROM nexus_clearing_transaction
                            WHERE batch_id = %s AND custody_state = 'PENDING'
                        )
                        UPDATE nexus_clearing_transaction transaction
                        SET selected = pending.command_rank <= %s,
                            reconciliation_state = CASE
                                WHEN pending.command_rank <= %s THEN 'STALE'
                                ELSE 'EXCLUDED'
                            END,
                            reconciliation_payload = '{}'::jsonb,
                            updated_at = now()
                        FROM pending
                        WHERE transaction.fingerprint = pending.fingerprint
                        """,
                        (execution["batch_id"], CLEARING_TRANCHE_SIZE, CLEARING_TRANCHE_SIZE),
                    )
                    progress = cursor.execute(
                        """
                        SELECT
                            COUNT(*) FILTER (WHERE custody_state = 'COMMITTED') AS committed_count,
                            COUNT(*) FILTER (WHERE custody_state = 'PENDING') AS remaining_count,
                            COUNT(*) FILTER (WHERE custody_state = 'PENDING' AND selected) AS next_armed_count
                        FROM nexus_clearing_transaction
                        WHERE batch_id = %s
                        """,
                        (execution["batch_id"],),
                    ).fetchone()
                    committed_count = int(progress["committed_count"] or 0)
                    remaining_count = int(progress["remaining_count"] or 0)
                    next_armed_count = int(progress["next_armed_count"] or 0)
                    batch_status = "IMPORTED" if remaining_count else "COMPLETED"
                    cursor.execute(
                        """
                        UPDATE nexus_clearing_batch
                        SET status = %s,
                            selected_count = %s,
                            reconciliation_summary = %s::jsonb,
                            approved_payload = NULL,
                            payload_hash = NULL,
                            submitted_by = NULL,
                            submitted_at = NULL,
                            approved_by = NULL,
                            approved_at = NULL,
                            updated_at = now()
                        WHERE batch_id = %s
                        """,
                        (
                            batch_status,
                            next_armed_count,
                            _json({
                                "committed": committed_count,
                                "remaining": remaining_count,
                                "next_armed": next_armed_count,
                            }),
                            execution["batch_id"],
                        ),
                    )
                else:
                    batch_status = {
                        "COMMIT_UNCERTAIN": "COMMIT_UNCERTAIN",
                        "BLOCKED": "BLOCKED",
                    }.get(status, "FAILED")
                    cursor.execute(
                        "UPDATE nexus_clearing_batch SET status = %s, updated_at = now() WHERE batch_id = %s",
                        (batch_status, execution["batch_id"]),
                    )
                self._audit_cursor(
                    cursor,
                    event_type="execution_committed" if status == "COMMITTED" else "execution_failed",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=execution["batch_id"],
                    reconciliation_id=(execution.get("approved_payload") or {}).get("reconciliation_id"),
                    execution_id=execution_id,
                    details={
                        "status": status,
                        "balance_mutations": len(evidence.get("balance_mutations", [])),
                        "queue_deletions": len(evidence.get("queue_deletions", [])),
                        "committed_count": committed_count,
                        "remaining_count": remaining_count,
                        "next_armed_count": next_armed_count,
                        "error": error,
                    },
                )
            connection.commit()
        return self.get_execution(execution_id)

    def finish_rollback(
        self,
        execution_id: str,
        *,
        actor: str,
        actor_role: str,
        reason: str,
        evidence: dict[str, Any],
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                execution = cursor.execute(
                    "SELECT batch_id FROM nexus_clearing_execution_run WHERE execution_id = %s FOR UPDATE",
                    (execution_id,),
                ).fetchone()
                if not execution:
                    raise LookupError("Clearing execution not found.")
                cursor.execute(
                    """
                    UPDATE nexus_clearing_execution_run
                    SET status = 'ROLLED_BACK',
                        rollback_requested_by = %s,
                        rollback_reason = %s,
                        rolled_back_at = now(),
                        oracle_evidence = oracle_evidence || %s::jsonb
                    WHERE execution_id = %s
                    """,
                    (actor, reason, _json({"rollback": evidence}), execution_id),
                )
                cursor.execute(
                    """
                    UPDATE nexus_clearing_execution_item
                    SET status = 'ROLLED_BACK', updated_at = now()
                    WHERE execution_id = %s
                    """,
                    (execution_id,),
                )
                cursor.execute(
                    "UPDATE nexus_clearing_batch SET status = 'ROLLED_BACK', updated_at = now() WHERE batch_id = %s",
                    (execution["batch_id"],),
                )
                self._audit_cursor(
                    cursor,
                    event_type="execution_rolled_back",
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=execution["batch_id"],
                    execution_id=execution_id,
                    details={"reason": reason, **evidence},
                )
            connection.commit()
        return self.get_execution(execution_id)

    def record_rollback_failure(
        self,
        execution_id: str,
        *,
        actor: str,
        actor_role: str,
        reason: str,
        error: str,
        commit_uncertain: bool,
    ) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                execution = cursor.execute(
                    """
                    SELECT batch_id, status
                    FROM nexus_clearing_execution_run
                    WHERE execution_id = %s
                    FOR UPDATE
                    """,
                    (execution_id,),
                ).fetchone()
                if not execution:
                    raise LookupError("Clearing execution not found.")
                if commit_uncertain:
                    cursor.execute(
                        """
                        UPDATE nexus_clearing_execution_run
                        SET status = 'COMMIT_UNCERTAIN',
                            error_message = %s,
                            rollback_requested_by = %s,
                            rollback_reason = %s
                        WHERE execution_id = %s
                        """,
                        (error[:2000], actor, reason, execution_id),
                    )
                    cursor.execute(
                        """
                        UPDATE nexus_clearing_batch
                        SET status = 'COMMIT_UNCERTAIN', updated_at = now()
                        WHERE batch_id = %s
                        """,
                        (execution["batch_id"],),
                    )
                self._audit_cursor(
                    cursor,
                    event_type=(
                        "rollback_commit_uncertain"
                        if commit_uncertain
                        else "rollback_blocked"
                    ),
                    actor=actor,
                    actor_role=actor_role,
                    batch_id=execution["batch_id"],
                    execution_id=execution_id,
                    details={
                        "reason": reason,
                        "error": error[:2000],
                        "commit_uncertain": commit_uncertain,
                    },
                )
            connection.commit()
        return self.get_execution(execution_id)

    def get_execution(self, execution_id: str) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                run = cursor.execute(
                    "SELECT * FROM nexus_clearing_execution_run WHERE execution_id = %s",
                    (execution_id,),
                ).fetchone()
                if not run:
                    raise LookupError("Clearing execution not found.")
                items = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_execution_item
                    WHERE execution_id = %s
                    ORDER BY updated_at, fingerprint
                    """,
                    (execution_id,),
                ).fetchall()
        return {**_plain(dict(run)), "items": [_plain(dict(item)) for item in items]}

    def list_executions(self, limit: int = 100) -> list[dict[str, Any]]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                rows = cursor.execute(
                    """
                    SELECT execution_id, batch_id, idempotency_key, payload_hash, status,
                           requested_by, reason, error_message, created_at, started_at,
                           completed_at, rollback_requested_by, rollback_reason, rolled_back_at
                    FROM nexus_clearing_execution_run
                    ORDER BY created_at DESC
                    LIMIT %s
                    """,
                    (max(1, min(limit, 500)),),
                ).fetchall()
        return [_plain(dict(row)) for row in rows]

    def list_audit(
        self,
        batch_id: str | None = None,
        limit: int = 150,
        search: str | None = None,
    ) -> list[dict[str, Any]]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                term = _normalized_text(search)
                pattern = f"%{term}%"
                rows = cursor.execute(
                    """
                    SELECT a.*
                    FROM nexus_clearing_audit a
                    WHERE (%s::text IS NULL OR a.batch_id = %s)
                      AND (
                          %s = ''
                          OR a.details::text ILIKE %s
                          OR COALESCE(a.batch_id, '') ILIKE %s
                          OR COALESCE(a.execution_id, '') ILIKE %s
                          OR EXISTS (
                              SELECT 1
                              FROM nexus_clearing_transaction t
                              WHERE t.batch_id = a.batch_id
                                AND (t.from_account ILIKE %s OR t.rrn ILIKE %s)
                          )
                      )
                    ORDER BY a.occurred_at DESC
                    LIMIT %s
                    """,
                    (
                        batch_id,
                        batch_id,
                        term,
                        pattern,
                        pattern,
                        pattern,
                        pattern,
                        pattern,
                        max(1, min(limit, 500)),
                    ),
                ).fetchall()
        return [_plain(dict(row)) for row in rows]

    def get_audit_evidence(self, audit_id: str) -> dict[str, Any]:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                event = cursor.execute(
                    "SELECT * FROM nexus_clearing_audit WHERE audit_id = %s",
                    (audit_id,),
                ).fetchone()
                if not event:
                    raise LookupError("Clearing audit event not found.")
                event = _plain(dict(event))
                approvals = []
                reconciliation = None
                snapshots = []
                execution_meta = None
                if event.get("execution_id"):
                    execution_meta = cursor.execute(
                        """
                        SELECT payload_hash, approved_payload, created_at
                        FROM nexus_clearing_execution_run
                        WHERE execution_id = %s
                        """,
                        (event["execution_id"],),
                    ).fetchone()
                    execution_meta = _plain(dict(execution_meta)) if execution_meta else None
                if event.get("batch_id"):
                    payload_hash = execution_meta.get("payload_hash") if execution_meta else None
                    if payload_hash:
                        approvals = cursor.execute(
                            """
                            SELECT * FROM nexus_clearing_approval
                            WHERE batch_id = %s AND payload_hash = %s
                            ORDER BY submitted_at DESC
                            """,
                            (event["batch_id"], payload_hash),
                        ).fetchall()
                    else:
                        approvals = cursor.execute(
                            """
                            SELECT * FROM nexus_clearing_approval
                            WHERE batch_id = %s
                            ORDER BY submitted_at DESC
                            """,
                            (event["batch_id"],),
                        ).fetchall()
                reconciliation_id = event.get("reconciliation_id")
                if not reconciliation_id and execution_meta:
                    reconciliation_id = (
                        execution_meta.get("approved_payload") or {}
                    ).get("reconciliation_id")
                if not reconciliation_id and event.get("batch_id"):
                    evidence_time = (
                        execution_meta.get("created_at")
                        if execution_meta
                        else event.get("occurred_at")
                    )
                    latest = cursor.execute(
                        """
                        SELECT reconciliation_id FROM nexus_clearing_reconciliation_run
                        WHERE batch_id = %s
                          AND completed_at IS NOT NULL
                          AND completed_at <= %s
                        ORDER BY completed_at DESC, started_at DESC
                        LIMIT 1
                        """,
                        (event["batch_id"], evidence_time),
                    ).fetchone()
                    reconciliation_id = latest["reconciliation_id"] if latest else None
                if reconciliation_id:
                    reconciliation_row = cursor.execute(
                        "SELECT * FROM nexus_clearing_reconciliation_run WHERE reconciliation_id = %s",
                        (reconciliation_id,),
                    ).fetchone()
                    reconciliation = _plain(dict(reconciliation_row)) if reconciliation_row else None
                    snapshots = cursor.execute(
                        """
                        SELECT * FROM nexus_clearing_account_snapshot
                        WHERE reconciliation_id = %s
                        ORDER BY external_account, account_role, captured_at
                        """,
                        (reconciliation_id,),
                    ).fetchall()

        batch = self.get_batch(event["batch_id"]) if event.get("batch_id") else None
        execution = self.get_execution(event["execution_id"]) if event.get("execution_id") else None
        sealed_intent = None
        if execution:
            sealed_intent = execution.get("approved_payload")
        elif batch:
            sealed_intent = batch.get("approved_payload")
        return {
            "event": event,
            "batch": batch,
            "sealed_intent": sealed_intent,
            "approvals": [_plain(dict(item)) for item in approvals],
            "reconciliation": reconciliation,
            "snapshots": [_plain(dict(item)) for item in snapshots],
            "execution": execution,
        }

    def record_account_lookup(
        self,
        external_account: str,
        *,
        actor: str,
        actor_role: str,
        evidence_summary: dict[str, Any],
    ) -> None:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                self._audit_cursor(
                    cursor,
                    event_type="live_account_inspected",
                    actor=actor,
                    actor_role=actor_role,
                    details={
                        "external_account": external_account,
                        **evidence_summary,
                    },
                )
            connection.commit()

    def record_rrn_lookup(
        self,
        rrn: str,
        *,
        actor: str,
        actor_role: str,
        evidence_summary: dict[str, Any],
    ) -> None:
        with self._connection() as connection:
            with connection.cursor() as cursor:
                self._audit_cursor(
                    cursor,
                    event_type="rrn_account_inspected",
                    actor=actor,
                    actor_role=actor_role,
                    details={"rrn": rrn, **evidence_summary},
                )
            connection.commit()

    @contextmanager
    def action_lock(self, key: str):
        if not self._dsn:
            lock = self._local_locks.setdefault(key, threading.Lock())
            if not lock.acquire(blocking=False):
                raise RuntimeError("A clearing operation for this batch is already in progress.")
            try:
                yield
            finally:
                lock.release()
            return
        connection = psycopg.connect(self._dsn, row_factory=dict_row)
        try:
            with connection.cursor() as cursor:
                row = cursor.execute(
                    "SELECT pg_try_advisory_lock(hashtext(%s)) AS locked",
                    (key,),
                ).fetchone()
                if not row or not row["locked"]:
                    raise RuntimeError("A clearing operation for this batch is already in progress.")
            yield
        finally:
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT pg_advisory_unlock(hashtext(%s))", (key,))
                connection.commit()
            finally:
                connection.close()

    @staticmethod
    def _notify_cursor(
        cursor,
        *,
        event_type: str,
        actor: str,
        batch_id: str | None = None,
        execution_id: str | None = None,
        fingerprint: str | None = None,
        audit_id: str | None = None,
        details: dict[str, Any] | None = None,
    ) -> None:
        cursor.execute(
            "SELECT pg_notify('nexus_clearing_events', %s)",
            (
                _json(
                    {
                        "type": "CUSTODY_UPDATE",
                        "event_type": event_type,
                        "audit_id": audit_id,
                        "batch_id": batch_id,
                        "execution_id": execution_id,
                        "fingerprint": fingerprint,
                        "actor": actor,
                        "details": details or {},
                    }
                ),
            ),
        )

    @staticmethod
    def _audit_cursor(
        cursor,
        *,
        event_type: str,
        actor: str,
        actor_role: str,
        batch_id: str | None = None,
        reconciliation_id: str | None = None,
        execution_id: str | None = None,
        fingerprint: str | None = None,
        details: dict[str, Any] | None = None,
    ) -> None:
        audit_id = f"nexus-clear-audit-{uuid4()}"
        cursor.execute(
            """
            INSERT INTO nexus_clearing_audit (
                audit_id, event_type, actor, actor_role, batch_id,
                reconciliation_id, execution_id, fingerprint, details
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
            """,
            (
                audit_id,
                event_type,
                actor,
                actor_role,
                batch_id,
                reconciliation_id,
                execution_id,
                fingerprint,
                _json(details or {}),
            ),
        )
        ClearingRepository._notify_cursor(
            cursor,
            event_type=event_type,
            actor=actor,
            batch_id=batch_id,
            execution_id=execution_id,
            fingerprint=fingerprint,
            audit_id=audit_id,
        )


def _oracle_day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    try:
        return _parse_date(value)
    except ValueError:
        return None


def _queue_matches(
    transaction: dict[str, Any],
    queue: dict[str, Any],
    *,
    strict_amount: bool = True,
) -> bool:
    currency_matches = not transaction.get("currency") or _queue_account_currency(queue) == _currency(
        transaction.get("currency")
    )
    row_matches = not transaction.get("queue_row_id") or queue.get("row_id") == transaction.get("queue_row_id")
    amount_matches = (
        _money(queue.get("amount", 0)) == _money(transaction.get("amount", 0))
        and _money(queue.get("fee", 0)) == _money(transaction.get("charge", 0))
    )
    return (
        str(queue.get("internal_account")) == str(transaction.get("from_internal_account"))
        and _normalized_text(queue.get("rrn")) == _normalized_text(transaction.get("rrn"))
        and _normalized_text(queue.get("stan")) == _normalized_text(transaction.get("stan"))
        and _oracle_day(queue.get("entry_date")) == _parse_date(transaction.get("entry_date"))
        and not _is_explicit_credit_queue(queue)
        and currency_matches
        and row_matches
        and (amount_matches or not strict_amount)
    )


def summarize_account_evidence(
    balance_rows: list[dict[str, Any]],
    queue_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    unauthorized_debit = sum(
        (_money(row.get("unauth_db_sum", 0)) for row in balance_rows),
        Decimal("0"),
    )
    debit_queue_rows = [row for row in queue_rows if not _is_explicit_credit_queue(row)]
    queued_amount = sum(
        (_money(row.get("amount", 0)) for row in debit_queue_rows),
        Decimal("0"),
    )
    queued_fees = sum(
        (
            sum(
                (_money(row.get(key, 0)) for key in ("fee", "processing_fee", "card_billing_fee")),
                Decimal("0"),
            )
            for row in debit_queue_rows
        ),
        Decimal("0"),
    )
    physical_queue_amount = sum(
        (_money(row.get("debit_queue_amount", 0)) for row in balance_rows),
        Decimal("0"),
    )
    return _plain(
        {
            "physical_row_count": len(balance_rows),
            "queue_row_count": len(debit_queue_rows),
            "unauthorized_debit_total": unauthorized_debit,
            "queued_amount_total": queued_amount,
            "queued_fee_total": queued_fees,
            "physical_queue_amount_total": physical_queue_amount,
            "orphaned_queue_amount_total": max(physical_queue_amount - queued_amount - queued_fees, Decimal("0")),
            "currencies": sorted({_currency(row.get("currency")) for row in balance_rows if row.get("currency")}),
        }
    )


def resolve_balance_rows(
    rows: list[dict[str, Any]],
    *,
    column: Literal[
        "ACNTBAL_AC_UNAUTH_DB_SUM",
        "ACNTBAL_AC_DB_QUEUE_AMT",
    ] = "ACNTBAL_AC_UNAUTH_DB_SUM",
    delta: Decimal,
    currency: str | None = None,
) -> dict[str, Any]:
    key = PHYSICAL_MUTATION_KEYS[column]
    requested_currency = _currency(currency)
    normalized = [
        {**row, key: _money(row.get(key, 0)), "currency": _currency(row.get("currency"))}
        for row in rows
        if not requested_currency or _currency(row.get("currency")) == requested_currency
    ]
    if not normalized:
        return {
            "state": "BLOCKED",
            "column": column,
            "delta": delta,
            "currency": requested_currency or None,
            "message": (
                f"No physical ACNTBAL row exists for currency {requested_currency}."
                if requested_currency
                else "No physical ACNTBAL row exists for the resolved internal account."
            ),
        }
    nonzero = [row for row in normalized if row[key] != 0]
    if not nonzero:
        return {
            "state": "ZERO",
            "column": column,
            "delta": delta,
            "current": Decimal("0"),
            "currency": requested_currency or None,
            "message": (
                "The physical debit queue amount is already zero."
                if column == "ACNTBAL_AC_DB_QUEUE_AMT"
                else "The relevant unauthorized balance component is already zero."
            ),
        }
    if len(nonzero) > 1:
        exact = [row for row in nonzero if row[key] == delta]
        reason = (
            "An exact component exists beside another non-zero physical row; formal resolver confirmation is required."
            if len(exact) == 1
            else "Multiple non-zero physical ACNTBAL rows can satisfy this account effect."
        )
        return {
            "state": "SPECIAL_REVIEW" if len(exact) == 1 else "BLOCKED",
            "column": column,
            "delta": delta,
            "candidate_count": len(nonzero),
            "currency": requested_currency or None,
            "message": reason,
        }
    target = nonzero[0]
    current = target[key]
    if current < delta:
        return {
            "state": "BLOCKED",
            "column": column,
            "delta": delta,
            "current": current,
            "row_id": target.get("row_id"),
            "currency": target.get("currency") or requested_currency or None,
            "message": (
                "The physical debit queue amount is smaller than the sealed queue repair."
                if column == "ACNTBAL_AC_DB_QUEUE_AMT"
                else "The resolved physical balance is smaller than the Finance-authorized effect."
            ),
        }
    after = current - delta
    return {
        "state": "SAFE",
        "column": column,
        "row_id": target["row_id"],
        "internal_account": str(target["internal_account"]),
        "currency": target.get("currency") or requested_currency or None,
        "current": current,
        "delta": delta,
        "after": after,
        "residual": after > 0,
        "physical_row_count": len(normalized),
        "message": (
            "One non-zero physical row is deterministically resolved; the unrelated residual will remain."
            if after > 0
            else "One non-zero physical row exactly carries the authorized effect."
        ),
    }


class ClearingOracleGateway:
    """Read and mutate only the exact Oracle rows approved by SentinelOps."""

    MAPPING_SQL = """
        SELECT ACNTS_INTERNAL_ACNUM AS internal_account
        FROM ACNTS
        WHERE FACNO(:entity_num, ACNTS_INTERNAL_ACNUM) = :external_account
        ORDER BY ACNTS_INTERNAL_ACNUM
    """
    REVERSE_ACCOUNT_SQL = """
        SELECT
            ACNTS_INTERNAL_ACNUM AS internal_account,
            FACNO(:entity_num, ACNTS_INTERNAL_ACNUM) AS external_account
        FROM ACNTS
        WHERE ACNTS_INTERNAL_ACNUM = :internal_account
        ORDER BY ACNTS_INTERNAL_ACNUM
    """
    ACCOUNT_SQL = """
        SELECT
            ACNTS_INTERNAL_ACNUM AS internal_account,
            ACNTS_ACCOUNT_NUMBER AS account_number,
            ACNTS_BRN_CODE AS branch_code,
            ACNTS_CLIENT_NUM AS client_number,
            ACNTS_PROD_CODE AS product_code,
            ACNTS_AC_TYPE AS account_type,
            ACNTS_AC_SUB_TYPE AS account_sub_type,
            ACNTS_AC_NAME1 AS account_name,
            ACNTS_SHORT_NAME AS short_name,
            ACNTS_CURR_CODE AS account_currency,
            ACNTS_CREATION_STATUS AS creation_status,
            ACNTS_DB_FREEZED AS debit_frozen,
            ACNTS_CR_FREEZED AS credit_frozen,
            ACNTS_DORMANT_ACNT AS dormant,
            ACNTS_LAST_TRAN_DATE AS last_transaction_date
        FROM ACNTS
        WHERE FACNO(:entity_num, ACNTS_INTERNAL_ACNUM) = :external_account
        ORDER BY ACNTS_INTERNAL_ACNUM
    """
    BALANCE_SQL = """
        SELECT
            ROWIDTOCHAR(ROWID) AS row_id,
            ACNTBAL_ENTITY_NUM AS entity_num,
            ACNTBAL_INTERNAL_ACNUM AS internal_account,
            TRIM(ACNTBAL_CURR_CODE) AS currency,
            NVL(ACNTBAL_AC_CUR_DB_SUM, 0) AS current_debit_sum,
            NVL(ACNTBAL_AC_CUR_CR_SUM, 0) AS current_credit_sum,
            NVL(ACNTBAL_AC_BAL, 0) AS account_balance,
            NVL(ACNTBAL_BC_BAL, 0) AS base_currency_balance,
            NVL(ACNTBAL_AC_UNAUTH_DB_SUM, 0) AS unauth_db_sum,
            NVL(ACNTBAL_AC_DB_QUEUE_AMT, 0) AS debit_queue_amount,
            NVL(ACNTBAL_AC_AMT_ON_HOLD, 0) AS amount_on_hold,
            NVL(ACNTBAL_AC_LIEN_AMT, 0) AS lien_amount,
            NVL(ACNTBAL_AC_CLG_DB_SUM, 0) AS clearing_debit_sum,
            ORA_ROWSCN AS row_scn
        FROM ACNTBAL
        WHERE ACNTBAL_INTERNAL_ACNUM = :internal_account
        ORDER BY ACNTBAL_CURR_CODE, ROWIDTOCHAR(ROWID)
    """
    QUEUE_SQL = """
        SELECT
            ROWIDTOCHAR(ROWID) AS row_id,
            BGPQ_ENTITY_NUM AS entity_num,
            BGPQ_ASIIN_BRN_CODE AS branch_code,
            BGPQ_ASIIN_SERIAL AS serial_number,
            BGPQ_INT_ACCT_1 AS internal_account,
            BGPQ_INT_ACCT_2 AS destination_internal_account,
            BGPQ_ASIIN_DATE AS entry_date,
            TRIM(BGPQ_REQ_RRN) AS rrn,
            TRIM(BGPQ_REQ_STAN_NO) AS stan,
            NVL(BGPQ_TRN_AMT, 0) AS amount,
            NVL(BGPQ_TRN_FEE, 0) AS fee,
            NVL(BGPQ_TRN_PROC_FEE, 0) AS processing_fee,
            NVL(BGPQ_CRD_BILL_FEE, 0) AS card_billing_fee,
            TRIM(BGPQ_CCY_CODE) AS currency,
            TRIM(BGPQ_ACT_CCY) AS account_currency,
            NVL(BGPQ_TRN_AMT_TCY, 0) AS transaction_currency_amount,
            TRIM(BGPQ_FLG_DRCR) AS debit_credit_flag,
            TRIM(BGPQ_REQ_PC) AS processing_code,
            TRIM(BGPQ_REQ_MTI) AS message_type,
            BGPQ_REQ_TIME AS request_time,
            TRIM(BGPQ_REQ_ACQ_ID) AS acquirer_id,
            TRIM(BGPQ_REQ_TRMID) AS terminal_id,
            TRIM(BGPQ_CHNL_ID) AS channel_id,
            TRIM(BGPQ_TRN_IND) AS transaction_indicator,
            TRIM(BGPQ_TRN_NARR1) AS narration_1,
            TRIM(BGPQ_TRN_NARR2) AS narration_2,
            TRIM(BGPQ_TRN_NARR3) AS narration_3,
            TRIM(BGPQ_TRN_CHGCD) AS charge_code
        FROM ASIBGPQNP
        WHERE BGPQ_INT_ACCT_1 = :internal_account
        ORDER BY BGPQ_ASIIN_DATE, ROWIDTOCHAR(ROWID)
    """
    RRN_LOOKUP_SQL = """
        SELECT
            ROWIDTOCHAR(ROWID) AS row_id,
            BGPQ_ENTITY_NUM AS entity_num,
            BGPQ_ASIIN_BRN_CODE AS branch_code,
            BGPQ_ASIIN_SERIAL AS serial_number,
            BGPQ_INT_ACCT_1 AS internal_account,
            BGPQ_INT_ACCT_2 AS destination_internal_account,
            BGPQ_ASIIN_DATE AS entry_date,
            TRIM(BGPQ_REQ_RRN) AS rrn,
            TRIM(BGPQ_REQ_STAN_NO) AS stan,
            NVL(BGPQ_TRN_AMT, 0) AS amount,
            NVL(BGPQ_TRN_FEE, 0) AS fee,
            NVL(BGPQ_TRN_PROC_FEE, 0) AS processing_fee,
            NVL(BGPQ_CRD_BILL_FEE, 0) AS card_billing_fee,
            TRIM(BGPQ_CCY_CODE) AS currency,
            TRIM(BGPQ_ACT_CCY) AS account_currency,
            NVL(BGPQ_TRN_AMT_TCY, 0) AS transaction_currency_amount,
            TRIM(BGPQ_FLG_DRCR) AS debit_credit_flag,
            TRIM(BGPQ_REQ_PC) AS processing_code,
            TRIM(BGPQ_REQ_MTI) AS message_type,
            BGPQ_REQ_TIME AS request_time,
            TRIM(BGPQ_REQ_ACQ_ID) AS acquirer_id,
            TRIM(BGPQ_REQ_TRMID) AS terminal_id,
            TRIM(BGPQ_CHNL_ID) AS channel_id,
            TRIM(BGPQ_TRN_IND) AS transaction_indicator,
            TRIM(BGPQ_TRN_NARR1) AS narration_1,
            TRIM(BGPQ_TRN_NARR2) AS narration_2,
            TRIM(BGPQ_TRN_NARR3) AS narration_3,
            TRIM(BGPQ_TRN_CHGCD) AS charge_code
        FROM ASIBGPQNP
        WHERE INSTR(
            UPPER(REPLACE(NVL(BGPQ_TRN_NARR1, ''), ' ', '')),
            :rrn_token
        ) > 0
        ORDER BY BGPQ_ASIIN_DATE DESC, ROWIDTOCHAR(ROWID)
    """

    def collect(
        self,
        external_accounts: Iterable[str],
        debit_external_accounts: Iterable[str],
        direct_mappings: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        connection = self._connect(read_only=True)
        try:
            mappings: dict[str, list[str]] = {
                _normalized_account(locator): [_normalized_account(internal)]
                for locator, internal in (direct_mappings or {}).items()
            }
            balances: dict[str, list[dict[str, Any]]] = {}
            queues: dict[str, list[dict[str, Any]]] = {}
            cursor = connection.cursor()
            try:
                for external_account in sorted(set(external_accounts)):
                    if external_account in mappings:
                        continue
                    cursor.execute(
                        self.MAPPING_SQL,
                        {
                            "entity_num": settings.NEXUS_CLEARING_ENTITY_NUMBER,
                            "external_account": external_account,
                        },
                    )
                    mappings[external_account] = [str(row[0]) for row in cursor.fetchall()]
                internal_accounts = sorted(
                    {values[0] for values in mappings.values() if len(values) == 1}
                )
                for internal_account in internal_accounts:
                    balances[internal_account] = self._fetch_dicts(
                        cursor,
                        self.BALANCE_SQL,
                        {"internal_account": internal_account},
                    )
                for external_account in sorted(set(debit_external_accounts)):
                    mapped = mappings.get(external_account, [])
                    if len(mapped) == 1:
                        queues[mapped[0]] = [
                            row
                            for row in self._fetch_dicts(
                                cursor,
                                self.QUEUE_SQL,
                                {"internal_account": mapped[0]},
                            )
                            if not _is_explicit_credit_queue(row)
                        ]
            finally:
                cursor.close()
            return {"mappings": mappings, "balances": balances, "queues": queues}
        finally:
            connection.close()

    def inspect_rrn(self, rrn: str) -> dict[str, Any]:
        lookup_rrn = _normalized_rrn_lookup(rrn)
        token = f"RRN:{lookup_rrn.upper()}"
        connection = self._connect(read_only=True)
        try:
            cursor = connection.cursor()
            try:
                matched_rows = [
                    row
                    for row in self._fetch_dicts(cursor, self.RRN_LOOKUP_SQL, {"rrn_token": token})
                    if not _is_explicit_credit_queue(row)
                ]
                internal_accounts = sorted(
                    {
                        str(row["internal_account"])
                        for row in matched_rows
                        if row.get("internal_account") not in (None, "")
                    }
                )
                balance_rows: list[dict[str, Any]] = []
                queue_rows: list[dict[str, Any]] = []
                external_accounts: list[str] = []
                if len(internal_accounts) == 1:
                    internal_account = internal_accounts[0]
                    external_accounts = sorted(
                        {
                            str(row["external_account"])
                            for row in self._fetch_dicts(
                                cursor,
                                self.REVERSE_ACCOUNT_SQL,
                                {
                                    "entity_num": settings.NEXUS_CLEARING_ENTITY_NUMBER,
                                    "internal_account": internal_account,
                                },
                            )
                            if row.get("external_account") not in (None, "")
                        }
                    )
                    balance_rows = self._fetch_dicts(
                        cursor,
                        self.BALANCE_SQL,
                        {"internal_account": internal_account},
                    )
                    matched_ids = {str(row.get("row_id")) for row in matched_rows}
                    queue_rows = [
                        {**row, "lookup_match": str(row.get("row_id")) in matched_ids}
                        for row in self._fetch_dicts(
                            cursor,
                            self.QUEUE_SQL,
                            {"internal_account": internal_account},
                        )
                        if not _is_explicit_credit_queue(row)
                    ]
                else:
                    queue_rows = [{**row, "lookup_match": True} for row in matched_rows]
            finally:
                cursor.close()
        finally:
            connection.close()
        return _plain(
            {
                "external_account": (
                    external_accounts[0]
                    if len(external_accounts) == 1
                    else internal_accounts[0]
                    if len(internal_accounts) == 1
                    else f"RRN-{lookup_rrn}"
                ),
                "external_accounts": external_accounts,
                "lookup_mode": "RRN",
                "lookup_value": lookup_rrn,
                "internal_accounts": internal_accounts,
                "account_profiles": [],
                "mapping_count": len(internal_accounts),
                "matched_queue_count": len(matched_rows),
                "balance_rows": balance_rows,
                "queue_rows": queue_rows,
                "summary": summarize_account_evidence(balance_rows, queue_rows),
                "captured_at": datetime.now(timezone.utc),
            }
        )

    def inspect_account(self, external_account: str) -> dict[str, Any]:
        account = _normalized_account(external_account)
        connection = self._connect(read_only=True)
        try:
            cursor = connection.cursor()
            try:
                cursor.execute(
                    self.ACCOUNT_SQL,
                    {
                        "entity_num": settings.NEXUS_CLEARING_ENTITY_NUMBER,
                        "external_account": account,
                    },
                )
                columns = [str(item[0]).lower() for item in cursor.description]
                account_profiles = [dict(zip(columns, row)) for row in cursor.fetchall()]
                internal_accounts = [str(row["internal_account"]) for row in account_profiles]
                balance_rows: list[dict[str, Any]] = []
                queue_rows: list[dict[str, Any]] = []
                for internal_account in internal_accounts:
                    balance_rows.extend(
                        self._fetch_dicts(
                            cursor,
                            self.BALANCE_SQL,
                            {"internal_account": internal_account},
                        )
                    )
                    queue_rows.extend(
                        row
                        for row in self._fetch_dicts(
                            cursor,
                            self.QUEUE_SQL,
                            {"internal_account": internal_account},
                        )
                        if not _is_explicit_credit_queue(row)
                    )
            finally:
                cursor.close()
        finally:
            connection.close()
        return _plain(
            {
                "external_account": account,
                "internal_accounts": internal_accounts,
                "account_profiles": account_profiles,
                "mapping_count": len(internal_accounts),
                "balance_rows": balance_rows,
                "queue_rows": queue_rows,
                "summary": summarize_account_evidence(balance_rows, queue_rows),
                "captured_at": datetime.now(timezone.utc),
            }
        )

    def execute_atomic(
        self,
        payload: dict[str, Any],
        *,
        execution_id: str,
        status_callback: Callable[[str], None] | None = None,
    ) -> dict[str, Any]:
        self._assert_writes_enabled()
        callback = status_callback or (lambda _status: None)
        connection = self._connect(read_only=False)
        evidence: dict[str, Any] = {
            "execution_id": execution_id,
            "payload_hash": payload["payload_hash"],
            "balance_mutations": [],
            "queue_deletions": [],
            "before": {},
            "after": {},
        }
        commit_started = False
        try:
            cursor = connection.cursor()
            try:
                callback("LOCKING")
                mutations = sorted(
                    payload["balance_mutations"],
                    key=lambda item: (item["row_id"], item["column"]),
                )
                queue_deletions = sorted(payload["queue_deletions"], key=lambda item: item["row_id"])
                locked_balances: dict[str, dict[str, Any]] = {}
                for row_id in sorted({item["row_id"] for item in mutations}):
                    cursor.execute(
                        """
                        SELECT
                            ROWIDTOCHAR(ROWID) AS row_id,
                            ACNTBAL_ENTITY_NUM AS entity_num,
                            ACNTBAL_INTERNAL_ACNUM AS internal_account,
                            TRIM(ACNTBAL_CURR_CODE) AS currency,
                            NVL(ACNTBAL_AC_UNAUTH_DB_SUM, 0) AS unauth_db_sum,
                            NVL(ACNTBAL_AC_DB_QUEUE_AMT, 0) AS debit_queue_amount
                        FROM ACNTBAL
                        WHERE ROWID = CHARTOROWID(:row_id)
                        FOR UPDATE NOWAIT
                        """,
                        {"row_id": row_id},
                    )
                    row = cursor.fetchone()
                    if not row:
                        raise RuntimeError(f"Approved ACNTBAL row {row_id} no longer exists.")
                    columns = [str(item[0]).lower() for item in cursor.description]
                    locked_balances[row_id] = dict(zip(columns, row))
                locked_queues: dict[str, dict[str, Any]] = {}
                for queue in queue_deletions:
                    cursor.execute(
                        """
                        SELECT
                            ROWIDTOCHAR(ROWID) AS row_id,
                            BGPQ_INT_ACCT_1 AS internal_account,
                            BGPQ_ASIIN_DATE AS entry_date,
                            TRIM(BGPQ_REQ_RRN) AS rrn,
                            TRIM(BGPQ_REQ_STAN_NO) AS stan,
                            NVL(BGPQ_TRN_AMT, 0) AS amount,
                            NVL(BGPQ_TRN_FEE, 0) AS fee,
                            NVL(BGPQ_TRN_PROC_FEE, 0) AS processing_fee,
                            NVL(BGPQ_CRD_BILL_FEE, 0) AS card_billing_fee,
                            TRIM(BGPQ_CCY_CODE) AS currency,
                            TRIM(BGPQ_ACT_CCY) AS account_currency,
                            TRIM(BGPQ_FLG_DRCR) AS debit_credit_flag
                        FROM ASIBGPQNP
                        WHERE ROWID = CHARTOROWID(:row_id)
                        FOR UPDATE NOWAIT
                        """,
                        {"row_id": queue["row_id"]},
                    )
                    row = cursor.fetchone()
                    if not row:
                        raise RuntimeError(f"Approved queue row {queue['row_id']} no longer exists.")
                    columns = [str(item[0]).lower() for item in cursor.description]
                    locked = dict(zip(columns, row))
                    if not _queue_matches(queue, locked):
                        raise RuntimeError(f"Queue row {queue['row_id']} changed after approval.")
                    locked_queues[queue["row_id"]] = locked

                callback("REVALIDATING")
                for mutation in mutations:
                    row = locked_balances[mutation["row_id"]]
                    column = mutation["column"]
                    if column not in PHYSICAL_MUTATION_KEYS:
                        raise RuntimeError("Execution payload contains a non-allowlisted physical column.")
                    current = _money(row[PHYSICAL_MUTATION_KEYS[column]])
                    if current != _money(mutation["before"]):
                        raise RuntimeError(
                            f"ACNTBAL row {mutation['row_id']} changed after approval "
                            f"({current} != {mutation['before']})."
                        )
                    if str(row["internal_account"]) != str(mutation["internal_account"]):
                        raise RuntimeError(f"ACNTBAL row {mutation['row_id']} changed account ownership.")
                    if _currency(row.get("currency")) != _currency(mutation.get("currency")):
                        raise RuntimeError(f"ACNTBAL row {mutation['row_id']} changed currency ownership.")

                queue_snapshots = {
                    row_id: self._capture_insertable_queue_row(cursor, row_id)
                    for row_id in locked_queues
                }
                callback("EXECUTING")
                for mutation in mutations:
                    column = mutation["column"]
                    if column not in PHYSICAL_MUTATION_KEYS:
                        raise RuntimeError("Execution payload contains a non-allowlisted physical column.")
                    cursor.execute(
                        f"""
                        UPDATE ACNTBAL
                        SET {column} = {column} - :delta
                        WHERE ROWID = CHARTOROWID(:row_id)
                          AND ACNTBAL_INTERNAL_ACNUM = :internal_account
                          AND TRIM(ACNTBAL_CURR_CODE) = :currency
                          AND NVL({column}, 0) = :expected
                        """,
                        {
                            "delta": _money(mutation["delta"]),
                            "row_id": mutation["row_id"],
                            "internal_account": mutation["internal_account"],
                            "currency": _currency(mutation["currency"]),
                            "expected": _money(mutation["before"]),
                        },
                    )
                    if cursor.rowcount != 1:
                        raise RuntimeError(f"Guarded update failed for ACNTBAL row {mutation['row_id']}.")
                    evidence["balance_mutations"].append(
                        {
                            **mutation,
                            "before": _money(mutation["before"]),
                            "delta": _money(mutation["delta"]),
                            "after": _money(mutation["after"]),
                        }
                    )
                for queue in queue_deletions:
                    cursor.execute(
                        "DELETE FROM ASIBGPQNP WHERE ROWID = CHARTOROWID(:row_id)",
                        {"row_id": queue["row_id"]},
                    )
                    if cursor.rowcount != 1:
                        raise RuntimeError(f"Guarded queue deletion failed for {queue['row_id']}.")
                    evidence["queue_deletions"].append(
                        {**queue, "insertable_row": queue_snapshots[queue["row_id"]]}
                    )

                callback("POST_VALIDATING")
                for mutation in mutations:
                    column = mutation["column"]
                    cursor.execute(
                        f"SELECT NVL({column}, 0) FROM ACNTBAL WHERE ROWID = CHARTOROWID(:row_id)",
                        {"row_id": mutation["row_id"]},
                    )
                    row = cursor.fetchone()
                    if not row or _money(row[0]) != _money(mutation["after"]):
                        raise RuntimeError(f"Post-validation failed for ACNTBAL row {mutation['row_id']}.")
                for queue in queue_deletions:
                    cursor.execute(
                        "SELECT COUNT(*) FROM ASIBGPQNP WHERE ROWID = CHARTOROWID(:row_id)",
                        {"row_id": queue["row_id"]},
                    )
                    if int(cursor.fetchone()[0]) != 0:
                        raise RuntimeError(f"Queue row {queue['row_id']} remains after guarded deletion.")

                evidence["before"] = {
                    "balance_rows": locked_balances,
                    "queue_rows": locked_queues,
                }
                evidence["after"] = {
                    "balance_mutations": evidence["balance_mutations"],
                    "deleted_queue_row_ids": [item["row_id"] for item in queue_deletions],
                }
                commit_started = True
                connection.commit()
                return _plain(evidence)
            finally:
                cursor.close()
        except Exception as exc:
            if not commit_started:
                try:
                    connection.rollback()
                except Exception:
                    pass
            if commit_started:
                raise CommitUncertainError(
                    "Oracle connection failed while commit outcome was being established. Do not retry blindly."
                ) from exc
            raise
        finally:
            connection.close()

    def rollback_atomic(
        self,
        execution_evidence: dict[str, Any],
        *,
        execution_id: str,
    ) -> dict[str, Any]:
        self._assert_writes_enabled()
        mutations = execution_evidence.get("balance_mutations", [])
        deletions = execution_evidence.get("queue_deletions", [])
        if not mutations and not deletions:
            raise RuntimeError("The committed execution does not contain complete compensating evidence.")
        connection = self._connect(read_only=False)
        result = {"execution_id": execution_id, "restored_balances": 0, "restored_queue_rows": 0}
        commit_started = False
        try:
            cursor = connection.cursor()
            try:
                for mutation in sorted(mutations, key=lambda item: (item["row_id"], item["column"])):
                    column = mutation["column"]
                    if column not in PHYSICAL_MUTATION_KEYS:
                        raise RuntimeError("Rollback evidence contains a non-allowlisted physical column.")
                    cursor.execute(
                        f"""
                        SELECT NVL({column}, 0)
                        FROM ACNTBAL
                        WHERE ROWID = CHARTOROWID(:row_id)
                          AND ACNTBAL_INTERNAL_ACNUM = :internal_account
                          AND TRIM(ACNTBAL_CURR_CODE) = :currency
                        FOR UPDATE NOWAIT
                        """,
                        {
                            "row_id": mutation["row_id"],
                            "internal_account": mutation["internal_account"],
                            "currency": _currency(mutation["currency"]),
                        },
                    )
                    row = cursor.fetchone()
                    if not row or _money(row[0]) != _money(mutation["after"]):
                        raise RuntimeError(
                            f"Rollback blocked because ACNTBAL row {mutation['row_id']} changed after clearing."
                        )
                    cursor.execute(
                        f"""
                        UPDATE ACNTBAL
                        SET {column} = {column} + :delta
                        WHERE ROWID = CHARTOROWID(:row_id)
                          AND ACNTBAL_INTERNAL_ACNUM = :internal_account
                          AND TRIM(ACNTBAL_CURR_CODE) = :currency
                          AND NVL({column}, 0) = :expected
                        """,
                        {
                            "delta": _money(mutation["delta"]),
                            "row_id": mutation["row_id"],
                            "internal_account": mutation["internal_account"],
                            "currency": _currency(mutation["currency"]),
                            "expected": _money(mutation["after"]),
                        },
                    )
                    if cursor.rowcount != 1:
                        raise RuntimeError(f"Rollback update failed for ACNTBAL row {mutation['row_id']}.")
                    result["restored_balances"] += 1
                for deletion in deletions:
                    cursor.execute(
                        """
                        SELECT COUNT(*)
                        FROM ASIBGPQNP
                        WHERE BGPQ_INT_ACCT_1 = :internal_account
                          AND TRIM(BGPQ_REQ_RRN) = :rrn
                          AND TRIM(BGPQ_REQ_STAN_NO) = :stan
                          AND BGPQ_ASIIN_DATE = :entry_date
                          AND ABS(NVL(BGPQ_TRN_AMT, 0) - :amount) <= 0.005
                          AND ABS(NVL(BGPQ_TRN_FEE, 0) - :fee) <= 0.005
                        """,
                        {
                            "internal_account": deletion["internal_account"],
                            "rrn": deletion["rrn"],
                            "stan": deletion["stan"],
                            "entry_date": _parse_date(deletion["entry_date"]),
                            "amount": _money(deletion["amount"]),
                            "fee": _money(deletion["charge"]),
                        },
                    )
                    if int(cursor.fetchone()[0]) != 0:
                        raise RuntimeError(
                            f"Rollback blocked because queue transaction {deletion['rrn']} already exists."
                        )
                    self._restore_queue_row(cursor, deletion["insertable_row"])
                    result["restored_queue_rows"] += 1
                commit_started = True
                connection.commit()
                return result
            finally:
                cursor.close()
        except Exception as exc:
            if not commit_started:
                try:
                    connection.rollback()
                except Exception:
                    pass
            if commit_started:
                raise CommitUncertainError(
                    "Oracle connection failed while rollback commit outcome was being established."
                ) from exc
            raise
        finally:
            connection.close()

    def _capture_insertable_queue_row(self, cursor, row_id: str) -> dict[str, Any]:
        cursor.execute(
            """
            SELECT COLUMN_NAME
            FROM USER_TAB_COLS
            WHERE TABLE_NAME = 'ASIBGPQNP'
              AND HIDDEN_COLUMN = 'NO'
              AND VIRTUAL_COLUMN = 'NO'
            ORDER BY COLUMN_ID
            """
        )
        columns = [str(row[0]).upper() for row in cursor.fetchall()]
        if not columns or any(not IDENTIFIER_RE.fullmatch(column) for column in columns):
            raise RuntimeError("Unable to establish a safe insertable ASIBGPQNP column set.")
        cursor.execute(
            f"SELECT {', '.join(columns)} FROM ASIBGPQNP WHERE ROWID = CHARTOROWID(:row_id)",
            {"row_id": row_id},
        )
        row = cursor.fetchone()
        if not row:
            raise RuntimeError(f"Queue row {row_id} disappeared before backup capture.")
        return {
            "columns": columns,
            "values": [_encode_oracle_value(value) for value in row],
        }

    @staticmethod
    def _restore_queue_row(cursor, snapshot: dict[str, Any]) -> None:
        columns = [str(column).upper() for column in snapshot.get("columns", [])]
        values = snapshot.get("values", [])
        if not columns or len(columns) != len(values):
            raise RuntimeError("Queue rollback snapshot is incomplete.")
        if any(not IDENTIFIER_RE.fullmatch(column) for column in columns):
            raise RuntimeError("Queue rollback snapshot contains an unsafe column identifier.")
        binds = {f"value_{index}": _decode_oracle_value(value) for index, value in enumerate(values)}
        placeholders = ", ".join(f":value_{index}" for index in range(len(values)))
        cursor.execute(
            f"INSERT INTO ASIBGPQNP ({', '.join(columns)}) VALUES ({placeholders})",
            binds,
        )
        if cursor.rowcount != 1:
            raise RuntimeError("Queue rollback insert did not restore exactly one row.")

    @staticmethod
    def _fetch_dicts(cursor, sql: str, binds: dict[str, Any]) -> list[dict[str, Any]]:
        cursor.execute(sql, binds)
        columns = [str(item[0]).lower() for item in cursor.description]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]

    def _connect(self, *, read_only: bool):
        try:
            import oracledb
        except ImportError as exc:
            raise RuntimeError("The oracledb package is required for unauthorized clearing.") from exc
        if read_only:
            username = settings.NEXUS_CLEARING_ORACLE_RO_USER or settings.ORACLE_USER
            password = settings.NEXUS_CLEARING_ORACLE_RO_PASSWORD or settings.ORACLE_PASSWORD
            dsn = settings.NEXUS_CLEARING_ORACLE_RO_DSN or settings.ORACLE_DSN
        else:
            username = settings.NEXUS_CLEARING_ORACLE_RW_USER
            password = settings.NEXUS_CLEARING_ORACLE_RW_PASSWORD
            dsn = settings.NEXUS_CLEARING_ORACLE_RW_DSN
        if not dsn and settings.IDC_ORACLE_HOST and (
            settings.ORACLE_SID or settings.ORACLE_SERVICE or settings.ORACLE_SERVICE_NAME
        ):
            port = int(settings.ORACLE_PORT or "1521")
            if settings.ORACLE_SID:
                dsn = oracledb.makedsn(settings.IDC_ORACLE_HOST, port, sid=settings.ORACLE_SID)
            else:
                dsn = oracledb.makedsn(
                    settings.IDC_ORACLE_HOST,
                    port,
                    service_name=settings.ORACLE_SERVICE_NAME or settings.ORACLE_SERVICE,
                )
        if not username or not password or not dsn:
            credential_type = "read-only" if read_only else "read/write"
            raise RuntimeError(f"Unauthorized clearing Oracle {credential_type} credentials are incomplete.")
        connection = oracledb.connect(
            user=username,
            password=password.get_secret_value(),
            dsn=dsn,
            config_dir=settings.NEXUS_CLEARING_ORACLE_CONFIG_DIR or None,
            tcp_connect_timeout=settings.ORACLE_CONNECT_TIMEOUT_SECONDS,
        )
        if read_only:
            cursor = connection.cursor()
            try:
                cursor.execute("SET TRANSACTION READ ONLY")
            finally:
                cursor.close()
        return connection

    @staticmethod
    def _assert_writes_enabled() -> None:
        if not settings.NEXUS_CLEARING_WRITES_ENABLED:
            raise RuntimeError("Unauthorized clearing writes are disabled by the Nexus safety gate.")
        if settings.is_production and not settings.NEXUS_CLEARING_PRODUCTION_WRITES_ENABLED:
            raise RuntimeError("Production unauthorized clearing writes require the separate production safety gate.")


class CommitUncertainError(RuntimeError):
    pass


def _encode_oracle_value(value: Any) -> dict[str, Any]:
    if value is None:
        return {"type": "null", "value": None}
    if hasattr(value, "read") and callable(value.read):
        materialized = value.read()
        if isinstance(materialized, bytes):
            return {
                "type": "lob_bytes",
                "value": base64.b64encode(materialized).decode("ascii"),
            }
        return {"type": "lob_text", "value": str(materialized)}
    if isinstance(value, Decimal):
        return {"type": "decimal", "value": format(value, "f")}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, bytes):
        return {"type": "bytes", "value": base64.b64encode(value).decode("ascii")}
    if isinstance(value, bool):
        return {"type": "bool", "value": value}
    if isinstance(value, int):
        return {"type": "int", "value": value}
    if isinstance(value, float):
        return {"type": "decimal", "value": str(value)}
    return {"type": "text", "value": str(value)}


def _decode_oracle_value(encoded: dict[str, Any]) -> Any:
    kind = encoded.get("type")
    value = encoded.get("value")
    if kind == "null":
        return None
    if kind == "decimal":
        return Decimal(str(value))
    if kind == "datetime":
        return datetime.fromisoformat(str(value))
    if kind == "date":
        return date.fromisoformat(str(value))
    if kind == "bytes":
        return base64.b64decode(str(value))
    if kind == "lob_bytes":
        return base64.b64decode(str(value))
    if kind == "lob_text":
        return str(value)
    if kind == "bool":
        return bool(value)
    if kind == "int":
        return int(value)
    return value


def reconcile_transactions(
    transactions: list[dict[str, Any]],
    oracle_state: dict[str, Any],
) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    mappings: dict[str, list[str]] = oracle_state["mappings"]
    balances: dict[str, list[dict[str, Any]]] = oracle_state["balances"]
    queues: dict[str, list[dict[str, Any]]] = oracle_state["queues"]

    physical_delta: dict[tuple[str, str, str], Decimal] = defaultdict(lambda: Decimal("0"))
    prepared: list[dict[str, Any]] = []
    exact_matches: dict[str, list[dict[str, Any]]] = {}
    for row in transactions:
        item = dict(row)
        item["amount"] = _money(item["amount"])
        item["charge"] = _money(item["charge"])
        item["clear_amount"] = _effective_debit(item)
        item["operation_mode"] = item.get("operation_mode") or "QUEUE_EXACT"
        item["from_internal_account"] = mappings.get(item["from_account"], [None])[0] if len(
            mappings.get(item["from_account"], [])
        ) == 1 else None
        candidates = []
        if item["operation_mode"] in QUEUE_LINKED_MODES and item.get("from_internal_account"):
            candidates = [
                queue
                for queue in queues.get(str(item["from_internal_account"]), [])
                if _queue_matches(
                    item,
                    queue,
                    strict_amount=item["operation_mode"] in {"QUEUE_EXACT", "QUEUE_ROW_ONLY"},
                )
            ]
        exact_matches[item["fingerprint"]] = candidates
        requested_currency = _currency(item.get("currency"))
        if (
            item["operation_mode"] == "BALANCE_ALL"
            and item["clear_amount"] <= 0
            and item.get("from_internal_account")
        ):
            physical_rows = balances.get(str(item["from_internal_account"]), [])
            positive_rows = [
                row
                for row in physical_rows
                if _money(row.get("unauth_db_sum", 0)) > 0
                and (not requested_currency or _currency(row.get("currency")) == requested_currency)
            ]
            if len(positive_rows) == 1:
                item["clear_amount"] = _money(positive_rows[0].get("unauth_db_sum", 0))
                item["amount"] = item["clear_amount"]
                requested_currency = _currency(positive_rows[0].get("currency"))
        queue_currency = _queue_account_currency(candidates[0]) if len(candidates) == 1 else ""
        item["currency"] = requested_currency or queue_currency
        item["currency_conflict"] = bool(requested_currency and queue_currency and requested_currency != queue_currency)
        if item["operation_mode"] == "QUEUE_ROW_ONLY":
            item["mutation_column"] = None
            item["mutation_delta"] = Decimal("0")
        elif item["operation_mode"] == "QUEUE_AMOUNT_RESET":
            item["mutation_column"] = "ACNTBAL_AC_DB_QUEUE_AMT"
            item["mutation_delta"] = item["amount"]
        else:
            item["mutation_column"] = "ACNTBAL_AC_UNAUTH_DB_SUM"
            item["mutation_delta"] = item["clear_amount"]
        if item["mutation_column"]:
            physical_delta[(item["from_account"], item["currency"], item["mutation_column"])] += item[
                "mutation_delta"
            ]
        prepared.append(item)

    resolutions: dict[tuple[str, str, str], dict[str, Any]] = {}
    for (external, currency, column), delta in physical_delta.items():
        mapped = mappings.get(external, [])
        resolutions[(external, currency, column)] = (
            resolve_balance_rows(
                balances.get(mapped[0], []),
                column=column,
                delta=delta,
                currency=currency or None,
            )
            if len(mapped) == 1
            else {
                "state": "BLOCKED",
                "column": column,
                "delta": delta,
                "currency": currency or None,
                "message": "No unique ACNTS mapping exists for the debit account.",
            }
        )

    matched_queue_ids: set[str] = set()
    for candidates in exact_matches.values():
        if len(candidates) == 1:
            matched_queue_ids.add(candidates[0]["row_id"])
    queue_assignment_counts = Counter(
        candidates[0]["row_id"]
        for candidates in exact_matches.values()
        if len(candidates) == 1
    )

    results: dict[str, dict[str, Any]] = {}
    for item in prepared:
        blockers: list[str] = []
        warnings: list[str] = []
        mutation_resolution = (
            resolutions[(item["from_account"], item["currency"], item["mutation_column"])]
            if item["mutation_column"]
            else {
                "state": "NOT_REQUIRED",
                "column": None,
                "delta": Decimal("0"),
                "currency": item["currency"],
                "message": "This custody action removes one exact queue row and does not mutate ACNTBAL.",
            }
        )
        queue_matches = exact_matches[item["fingerprint"]]
        if len(mappings.get(item["from_account"], [])) != 1:
            blockers.append("Debit account mapping is missing or ambiguous.")
        if not item["currency"]:
            blockers.append("The transaction currency cannot be established from Finance or the live queue row.")
        if item["currency_conflict"]:
            blockers.append("Finance currency does not match the exact live queue currency.")
        if len(queue_matches) > 1:
            blockers.append("More than one exact live queue row matches the Finance transaction.")
        if len(queue_matches) == 1 and queue_assignment_counts[queue_matches[0]["row_id"]] > 1:
            blockers.append(
                "The same live queue row matches more than one Finance transaction; "
                "the source must be resolved before execution."
            )
        if len(queue_matches) == 0 and item["operation_mode"] in QUEUE_LINKED_MODES:
            warnings.append("The exact Finance-authorized queue row is not present.")
        if mutation_resolution["state"] == "BLOCKED":
            blockers.append(mutation_resolution["message"])
        elif mutation_resolution["state"] == "SPECIAL_REVIEW":
            warnings.append(mutation_resolution["message"])
        if item["operation_mode"] == "QUEUE_PARTIAL" and len(queue_matches) == 1:
            queue = queue_matches[0]
            queue_authority = sum(
                (_money(queue.get(key, 0)) for key in ("amount", "fee")),
                Decimal("0"),
            )
            if item["clear_amount"] > queue_authority:
                blockers.append("The requested partial clear exceeds the selected queue transaction value.")

        from_internal = str(item.get("from_internal_account") or "")
        all_account_queues = queues.get(from_internal, [])
        currency_queues = [
            queue
            for queue in all_account_queues
            if _queue_account_currency(queue) == item["currency"] and not _is_explicit_credit_queue(queue)
        ]
        if item["operation_mode"] == "QUEUE_AMOUNT_RESET" and currency_queues:
            blockers.append(
                "Live queue rows still exist in this currency. Remove the exact rows instead of resetting the physical queue amount."
            )
        extra_count = sum(1 for queue in all_account_queues if queue["row_id"] not in matched_queue_ids)
        if extra_count:
            warnings.append(
                f"{extra_count} unapproved queue row(s) belong to this debit account and will remain untouched."
            )

        if blockers:
            state = "BLOCKED"
            explanation = "Execution is blocked because the live Oracle evidence is not uniquely resolvable."
        elif not queue_matches and item["operation_mode"] in QUEUE_LINKED_MODES:
            if item["operation_mode"] == "QUEUE_ROW_ONLY":
                state = "NO_LONGER_OUTSTANDING"
                explanation = "The sealed queue row is no longer present. No queue-only action remains."
            elif mutation_resolution["state"] == "ZERO":
                state = "NO_LONGER_OUTSTANDING"
                explanation = "The exact queue row is absent and the unauthorized debit is already zero."
            else:
                state = "BLOCKED"
                blockers.append("A balance effect remains without the exact authorized queue row.")
                explanation = "SentinelOps cannot prove a complete transaction effect without the exact queue evidence."
        elif item["operation_mode"] == "QUEUE_ROW_ONLY":
            state = "READY"
            explanation = "One exact debit queue row is proven; ACNTBAL unauthorized debit will remain untouched."
        elif item["operation_mode"] == "QUEUE_AMOUNT_RESET" and mutation_resolution["state"] == "ZERO":
            state = "NO_LONGER_OUTSTANDING"
            explanation = "The physical debit queue amount is already zero. No repair remains."
        elif item["operation_mode"] == "QUEUE_AMOUNT_RESET" and mutation_resolution["state"] == "SAFE":
            state = "READY"
            explanation = "An orphaned physical queue amount is proven with no live queue row in this currency."
        elif mutation_resolution["state"] == "ZERO" and item["operation_mode"] == "BALANCE_ALL":
            state = "NO_LONGER_OUTSTANDING"
            explanation = "This physical currency row has no unauthorized debit remaining. No clearing action is available."
        elif mutation_resolution["state"] in {"SPECIAL_REVIEW", "ZERO"}:
            state = "SPECIAL_REVIEW"
            explanation = (
                "The unauthorized debit row is zero or physically ambiguous. "
                "A documented special-case decision is required."
            )
        else:
            residual = bool(mutation_resolution.get("residual"))
            if residual:
                state = "READY_WITH_RESIDUAL"
                explanation = "The exact authorized effect is safe; unrelated residual balance will be preserved."
            elif extra_count:
                state = "SAFE_EXTRA_QUEUE"
                explanation = "The exact authorized transaction is safe; additional queue rows will be preserved."
            else:
                state = "READY"
                explanation = "Mapping, physical balance rows, and the exact queue row reconcile cleanly."

        queue_match = None
        if len(queue_matches) == 1:
            queue = queue_matches[0]
            queue_match = {
                "row_id": queue["row_id"],
                "internal_account": str(queue["internal_account"]),
                "entry_date": _oracle_day(queue["entry_date"]),
                "rrn": _normalized_text(queue["rrn"]),
                "stan": _normalized_text(queue["stan"]),
                "amount": _money(queue["amount"]),
                "charge": _money(queue["fee"]),
                "processing_fee": _money(queue.get("processing_fee", 0)),
                "card_billing_fee": _money(queue.get("card_billing_fee", 0)),
                "currency": _queue_account_currency(queue),
                "transaction_currency": _currency(queue.get("currency")),
                "account_currency": _queue_account_currency(queue),
                "destination_internal_account": str(queue.get("destination_internal_account") or ""),
                "branch_code": queue.get("branch_code"),
                "serial_number": queue.get("serial_number"),
                "debit_credit_flag": _normalized_text(queue.get("debit_credit_flag")),
                "processing_code": _normalized_text(queue.get("processing_code")),
                "message_type": _normalized_text(queue.get("message_type")),
                "terminal_id": _normalized_text(queue.get("terminal_id")),
                "channel_id": _normalized_text(queue.get("channel_id")),
                "narration": " ".join(
                    _normalized_text(queue.get(key))
                    for key in ("narration_1", "narration_2", "narration_3")
                    if _normalized_text(queue.get(key))
                ),
            }
        results[item["fingerprint"]] = _plain(
            {
                "fingerprint": item["fingerprint"],
                "state": state,
                "explanation": explanation,
                "blockers": blockers,
                "warnings": warnings,
                "from_mapping": {
                    "external_account": item["from_account"],
                    "matches": mappings.get(item["from_account"], []),
                    "internal_account": item.get("from_internal_account"),
                    "locator_type": (
                        "ASIBGPQNP_RRN"
                        if item.get("source_internal_account")
                        else "ACNTS_ACCOUNT"
                    ),
                },
                "destination_context": item.get("to_account") or None,
                "currency": item["currency"],
                "clear_amount": item["clear_amount"],
                "operation_mode": item["operation_mode"],
                "debit_resolution": mutation_resolution,
                "queue_match": queue_match,
                "delete_queue": item["operation_mode"] in {"QUEUE_EXACT", "QUEUE_ROW_ONLY"} and queue_match is not None,
                "extra_queue_count": extra_count,
            }
        )

    snapshots: list[dict[str, Any]] = []
    for external in sorted({item["from_account"] for item in prepared}):
        mapped = mappings.get(external, [])
        internal = mapped[0] if len(mapped) == 1 else None
        snapshots.append(
            _plain(
                {
                    "external_account": external,
                    "internal_account": internal,
                    "account_role": "DEBIT",
                    "mapping_count": len(mapped),
                    "balance_rows": balances.get(str(internal), []) if internal else [],
                    "queue_rows": queues.get(str(internal), []) if internal else [],
                    "resolution": {
                        f"{currency or 'UNRESOLVED'}:{column}": resolution
                        for (account, currency, column), resolution in resolutions.items()
                        if account == external
                    },
                }
            )
        )
    counts = Counter(result["state"].lower() for result in results.values())
    summary = {
        "selected": len(results),
        "ready": counts["ready"],
        "ready_with_residual": counts["ready_with_residual"],
        "safe_extra_queue": counts["safe_extra_queue"],
        "special_review": counts["special_review"],
        "no_longer_outstanding": counts["no_longer_outstanding"],
        "blocked": counts["blocked"],
        "captured_at": datetime.now(timezone.utc),
    }
    return results, snapshots, _plain(summary)


def build_execution_payload(batch: dict[str, Any], change_reference: str, note: str) -> dict[str, Any]:
    transactions = [item for item in batch["transactions"] if item["selected"]]
    if not transactions:
        raise ValueError("Select at least one safe Finance transaction.")
    invalid = [
        item
        for item in transactions
        if item["reconciliation_state"] not in SAFE_STATES
    ]
    if invalid:
        raise ValueError(
            "Only safely reconciled transactions can be submitted. Deselect or resolve review and blocked cases."
        )
    mutations: dict[tuple[str, str], dict[str, Any]] = {}
    queue_deletions: list[dict[str, Any]] = []
    transaction_payloads: list[dict[str, Any]] = []
    for item in transactions:
        evidence = item["reconciliation_payload"]
        resolution = evidence["debit_resolution"]
        if resolution.get("state") != "NOT_REQUIRED":
            if resolution.get("column") not in PHYSICAL_MUTATION_KEYS or not resolution.get("row_id"):
                raise ValueError("Reconciliation evidence does not identify one allowlisted physical mutation.")
            mutation_key = (resolution["row_id"], resolution["column"])
            candidate = {
                "role": "QUEUE" if resolution["column"] == "ACNTBAL_AC_DB_QUEUE_AMT" else "DEBIT",
                "external_account": item["from_account"],
                "internal_account": resolution["internal_account"],
                "currency": resolution["currency"],
                "row_id": resolution["row_id"],
                "column": resolution["column"],
                "before": resolution["current"],
                "delta": resolution["delta"],
                "after": resolution["after"],
            }
            if mutation_key in mutations and mutations[mutation_key] != candidate:
                raise ValueError("Reconciliation evidence contains conflicting physical balance mutations.")
            mutations[mutation_key] = candidate
        queue = evidence.get("queue_match")
        if not queue and item.get("operation_mode") in QUEUE_LINKED_MODES:
            raise ValueError(f"Transaction {item['rrn']} has no exact queue evidence.")
        if queue and evidence.get("delete_queue"):
            queue_deletions.append(
                {
                    **queue,
                    "fingerprint": item["fingerprint"],
                    "from_internal_account": queue["internal_account"],
                    "entry_date": item["entry_date"],
                    "amount": item["amount"],
                    "charge": item["charge"],
                }
            )
        transaction_payloads.append(
            {
                "fingerprint": item["fingerprint"],
                "source_row": item["source_row"],
                "entry_date": item["entry_date"],
                "from_account": item["from_account"],
                "source_internal_account": item.get("source_internal_account"),
                "to_account": item.get("to_account") or "",
                "currency": evidence.get("currency") or item.get("currency"),
                "rrn": item["rrn"],
                "stan": item["stan"],
                "amount": item["amount"],
                "charge": item["charge"],
                "clear_amount": (
                    evidence.get("clear_amount")
                    if evidence.get("clear_amount") is not None
                    else _effective_debit(item)
                ),
                "operation_mode": item.get("operation_mode") or "QUEUE_EXACT",
                "classification": item["reconciliation_state"],
            }
        )
    payload = _plain(
        {
            "batch_id": batch["batch_id"],
            "reconciliation_id": (batch.get("latest_reconciliation") or {}).get("reconciliation_id"),
            "finance_reference": batch["finance_reference"],
            "change_reference": change_reference.strip(),
            "submission_note": note.strip(),
            "source_sha256": batch["source_sha256"],
            "transactions": transaction_payloads,
            "balance_mutations": sorted(
                mutations.values(),
                key=lambda item: (item["internal_account"], item["column"], item["row_id"]),
            ),
            "queue_deletions": sorted(queue_deletions, key=lambda item: item["fingerprint"]),
        }
    )
    if not payload["balance_mutations"] and not payload["queue_deletions"]:
        raise ValueError("The sealed selection does not contain an executable physical change.")
    payload_hash = hashlib.sha256(_json(payload).encode("utf-8")).hexdigest()
    payload["payload_hash"] = payload_hash
    return payload


class UnauthorizedClearingService:
    """Application boundary for import, reconciliation, custody, and execution."""

    def __init__(
        self,
        repository: ClearingRepository | None = None,
        oracle: ClearingOracleGateway | None = None,
        clock: Callable[[], datetime] | None = None,
    ) -> None:
        self.repository = repository or ClearingRepository()
        self.oracle = oracle or ClearingOracleGateway()
        self.clock = clock or (lambda: datetime.now(timezone.utc))

    def _is_admin_role(self, actor_role: str) -> bool:
        return actor_role.strip().lower() in {role.lower() for role in settings.NEXUS_ADMIN_ROLES}

    def specific_record_policy(self, *, actor_role: str) -> dict[str, Any]:
        policy = self.repository.get_specific_record_policy()
        is_admin = self._is_admin_role(actor_role)
        return {
            **policy,
            "is_admin": is_admin,
            "can_specific_record_clear": is_admin or bool(policy["enabled"]),
        }

    def set_specific_record_policy(
        self,
        request: ClearingPolicyUpdateRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        return {
            **self.repository.set_specific_record_policy(
                request.enabled,
                actor=actor,
                actor_role=actor_role,
            ),
            "is_admin": True,
            "can_specific_record_clear": True,
        }

    def _assert_specific_record_allowed(self, *, actor_role: str) -> None:
        if self._is_admin_role(actor_role):
            return
        if not self.repository.get_specific_record_policy()["enabled"]:
            raise ValueError(
                "Specific-record clearing is disabled for operators. Clear all unauthorized debits or ask an administrator to enable it."
            )

    def overview(self, *, actor_role: str = "unknown") -> dict[str, Any]:
        policy = self.specific_record_policy(actor_role=actor_role)
        return {
            **self.repository.overview(),
            "operating_window": clearing_operating_window(self.clock()),
            "specific_record_policy": policy,
        }

    def _assert_batch_window_open(self, batch_id: str) -> dict[str, Any]:
        batch = self.repository.get_batch(batch_id)
        window = clearing_operating_window(self.clock())
        armed_count = sum(
            1
            for transaction in batch.get("transactions", [])
            if transaction.get("selected") and transaction.get("custody_state", "PENDING") == "PENDING"
        )
        if window["blocked"] and armed_count > window["batch_threshold"]:
            raise ValueError(
                f"The active command tranche exceeds {window['batch_threshold']} records. "
                "Exclude rows until the tranche is within the guarded limit."
            )
        return batch

    def _identity_transactions(
        self,
        identities: list[str],
        identity_kind: Literal["ACCOUNT", "RRN"],
    ) -> list[ClearingTransactionInput]:
        transactions: list[ClearingTransactionInput] = []
        now = self.clock().astimezone(ZoneInfo(settings.NEXUS_CLEARING_TIMEZONE))
        next_source_row = 2

        if identity_kind == "ACCOUNT":
            try:
                account_state = self.oracle.collect(set(identities), set())
                collection_error = ""
            except Exception as exc:
                logger.warning("Account identity intake will be sealed for later verification: %s", exc)
                account_state = {"mappings": {}, "balances": {}, "queues": {}}
                collection_error = str(exc)

            for identity in identities:
                mapped = account_state.get("mappings", {}).get(identity, [])
                internal_account = str(mapped[0]) if len(mapped) == 1 else None
                balance_rows = account_state.get("balances", {}).get(internal_account, []) if internal_account else []
                rows_by_currency: dict[str, list[dict[str, Any]]] = defaultdict(list)
                for row in balance_rows:
                    rows_by_currency[_currency(row.get("currency"))].append(row)
                positive_rows = [
                    row
                    for row in balance_rows
                    if _money(row.get("unauth_db_sum", 0)) > 0
                ]
                candidates = positive_rows or [
                    rows[0]
                    for currency, rows in sorted(rows_by_currency.items())
                    if currency and len(rows) == 1
                ]
                if not candidates:
                    candidates = [{}]

                for row in candidates:
                    currency = _currency(row.get("currency")) or None
                    row_group = rows_by_currency.get(currency or "", [])
                    ambiguous = bool(currency and len(row_group) != 1)
                    current = _money(row.get("unauth_db_sum", 0)) if row else Decimal("0")
                    clear_amount = current if current > 0 and not ambiguous else Decimal("0")
                    if collection_error:
                        posture = "Oracle verification pending; source retained after a transient read failure"
                    elif len(mapped) != 1:
                        posture = "Account mapping unresolved; verify Oracle evidence"
                    elif ambiguous:
                        posture = "Physical currency row is ambiguous; inspect before submission"
                    elif clear_amount <= 0:
                        posture = "No positive unauthorized debit at intake"
                    else:
                        posture = "All unauthorized debits in the physical currency row"
                    token = hashlib.sha256(
                        f"ACCOUNT:{identity}:{internal_account or 'UNRESOLVED'}:{currency or 'UNRESOLVED'}:{next_source_row}".encode("utf-8")
                    ).hexdigest().upper()
                    transactions.append(
                        ClearingTransactionInput(
                            source_row=next_source_row,
                            entry_date=now.date(),
                            from_account=identity,
                            source_internal_account=internal_account,
                            currency=currency,
                            rrn=f"BAL-{token[:12]}",
                            stan=token[-6:],
                            amount=clear_amount,
                            clear_amount=clear_amount,
                            operation_mode="BALANCE_ALL",
                            narration=f"Identity import / {posture}",
                        )
                    )
                    next_source_row += 1
            return transactions

        failures: list[str] = []
        for identity in identities:
            try:
                    evidence = self.oracle.inspect_rrn(identity)
                    internal_accounts = [str(value) for value in evidence.get("internal_accounts", [])]
                    matches = [
                        row for row in evidence.get("queue_rows", [])
                        if row.get("lookup_match") and not _is_explicit_credit_queue(row)
                    ]
                    if len(internal_accounts) != 1 or len(matches) != 1:
                        raise ValueError("RRN must resolve to exactly one debit queue row")
                    queue = matches[0]
                    account = str(evidence.get("external_account") or internal_accounts[0])
                    amount = _money(queue.get("amount", 0))
                    charge = _money(queue.get("fee", 0))
                    transactions.append(
                        ClearingTransactionInput(
                            source_row=next_source_row,
                            entry_date=_oracle_day(queue.get("entry_date")) or now.date(),
                            from_account=account,
                            source_internal_account=internal_accounts[0],
                            currency=_queue_account_currency(queue),
                            rrn=_normalized_reference(queue.get("rrn") or identity, "RRN"),
                            stan=_normalized_reference(queue.get("stan"), "STAN"),
                            amount=amount,
                            charge=charge,
                            clear_amount=amount + charge,
                            operation_mode="QUEUE_EXACT",
                            queue_row_id=str(queue.get("row_id")),
                            narration="Identity import / exact RRN debit",
                        )
                    )
                    next_source_row += 1
            except Exception as exc:
                failures.append(f"Row {next_source_row} ({identity}): {exc}")
        if failures:
            preview = "; ".join(failures[:8])
            suffix = f" and {len(failures) - 8} more" if len(failures) > 8 else ""
            raise ValueError(f"Oracle identity resolution failed: {preview}{suffix}.")
        return transactions

    def import_source(
        self,
        request: ClearingSourceImportRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        try:
            content = base64.b64decode(request.content_base64, validate=True)
        except Exception as exc:
            raise ValueError("The uploaded Finance source is not valid base64 content.") from exc
        if len(content) > settings.NEXUS_CLEARING_MAX_SOURCE_BYTES:
            max_megabytes = settings.NEXUS_CLEARING_MAX_SOURCE_BYTES / (1024 * 1024)
            raise ValueError(f"The Finance source exceeds the {max_megabytes:g} MB import limit.")
        identities = parse_identity_source(request.filename, content, request.identity_kind)
        if request.identity_kind == "RRN":
            self._assert_specific_record_allowed(actor_role=actor_role)
        transactions = self._identity_transactions(identities, request.identity_kind)
        source_hash = hashlib.sha256(content).hexdigest()
        return self.repository.create_batch(
            ClearingBatchCreateRequest(
                batch_name=request.batch_name,
                finance_reference=request.finance_reference,
                source_filename=request.filename,
                source_kind="IDENTITY_IMPORT",
                transactions=transactions,
            ),
            actor=actor,
            actor_role=actor_role,
            source_sha256=source_hash,
        )

    def create_batch(
        self,
        request: ClearingBatchCreateRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        return self.repository.create_batch(request, actor=actor, actor_role=actor_role)

    def _create_account_batch_from_evidence(
        self,
        account_locator: str,
        evidence: dict[str, Any],
        request: ClearingAccountBatchRequest,
        *,
        actor: str,
        actor_role: str,
        source_internal_account: str | None = None,
        source_filename: str = "account-explorer.json",
    ) -> dict[str, Any]:
        account = _normalized_account(account_locator)
        if evidence["mapping_count"] != 1:
            raise ValueError("Direct clearing requires one uniquely resolved internal account.")
        if source_internal_account:
            resolved = [str(value) for value in evidence.get("internal_accounts", [])]
            if resolved != [source_internal_account]:
                raise ValueError("The RRN no longer resolves to the expected internal account.")
        balance = next(
            (
                row
                for row in evidence["balance_rows"]
                if row.get("row_id") == request.balance_row_id
                and _currency(row.get("currency")) == request.currency
            ),
            None,
        )
        if not balance:
            raise ValueError("The selected ACNTBAL currency row no longer exists for this account.")
        available = _money(balance.get("unauth_db_sum", 0))
        physical_queue_amount = _money(balance.get("debit_queue_amount", 0))
        currency_queue_rows = [
            row
            for row in evidence.get("queue_rows", [])
            if str(row.get("internal_account")) == str(balance.get("internal_account"))
            and _queue_account_currency(row) == request.currency
            and not _is_explicit_credit_queue(row)
        ]
        if request.mode in {"QUEUE_TRANSACTION", "ALL_UNAUTHORIZED_DEBITS"} and available <= 0:
            raise ValueError("The selected currency row has no unauthorized debit to clear.")

        now = self.clock().astimezone(ZoneInfo(settings.NEXUS_CLEARING_TIMEZONE))
        transactions: list[ClearingTransactionInput] = []
        if request.mode == "ALL_UNAUTHORIZED_DEBITS":
            clear_amount = available
            operation_mode = "BALANCE_ALL"
            identity = uuid4().hex[:12].upper()
            rrn = f"BAL-{identity}"
            stan = identity[-6:]
            amount = clear_amount
            charge = Decimal("0")
            transactions.append(
                ClearingTransactionInput(
                    source_row=1,
                    entry_date=now.date(),
                    from_account=account,
                    source_internal_account=source_internal_account,
                    to_account="",
                    currency=request.currency,
                    rrn=rrn,
                    stan=stan,
                    amount=amount,
                    charge=charge,
                    clear_amount=clear_amount,
                    operation_mode=operation_mode,
                    narration=request.note or "Account Explorer authorized debit clearing",
                )
            )
        elif request.mode == "QUEUE_AMOUNT_RESET":
            if available > 0:
                raise ValueError(
                    "The selected row still has an unauthorized debit. Use the normal debit-clearing path first."
                )
            if physical_queue_amount <= 0:
                raise ValueError("The selected currency row has no physical debit queue amount to repair.")
            if currency_queue_rows:
                raise ValueError(
                    "Live queue rows still exist for this currency. Remove an exact queue row instead of resetting the aggregate."
                )
            clear_amount = Decimal("0")
            operation_mode = "QUEUE_AMOUNT_RESET"
            identity = uuid4().hex[:12].upper()
            rrn = f"QAM-{identity}"
            stan = identity[-6:]
            amount = physical_queue_amount
            charge = Decimal("0")
            transactions.append(
                ClearingTransactionInput(
                    source_row=1,
                    entry_date=now.date(),
                    from_account=account,
                    source_internal_account=source_internal_account,
                    to_account="",
                    currency=request.currency,
                    rrn=rrn,
                    stan=stan,
                    amount=amount,
                    charge=charge,
                    clear_amount=clear_amount,
                    operation_mode=operation_mode,
                    narration=request.note or "Account Explorer queue-only custody repair",
                )
            )
        else:
            requested_row_ids = request.queue_row_ids
            queues_by_id = {
                _normalized_text(row.get("row_id")): row
                for row in evidence["queue_rows"]
                if str(row.get("internal_account")) == str(balance.get("internal_account"))
                and _queue_account_currency(row) == request.currency
                and not _is_explicit_credit_queue(row)
            }
            selected_queues = [queues_by_id[row_id] for row_id in requested_row_ids if row_id in queues_by_id]
            if len(selected_queues) != len(requested_row_ids):
                raise ValueError("One or more selected queue transactions no longer exist in this account and currency.")
            queue_authorities = [
                sum((_money(queue.get(key, 0)) for key in ("amount", "fee")), Decimal("0"))
                for queue in selected_queues
            ]
            if request.mode == "QUEUE_ROW_ONLY":
                if available > 0:
                    raise ValueError(
                        "The selected row still has an unauthorized debit. Use the normal debit-clearing path first."
                    )
                requested_total = Decimal("0")
            else:
                aggregate_authority = sum(queue_authorities, Decimal("0"))
                if len(selected_queues) > 1:
                    if request.clear_amount is not None and _money(request.clear_amount) != aggregate_authority:
                        raise ValueError(
                            "A multi-row queue selection must clear the complete proven value of every selected row."
                        )
                    requested_total = aggregate_authority
                else:
                    if request.clear_amount is None:
                        raise ValueError("Enter the exact debit amount to clear for the selected queue transaction.")
                    requested_total = _money(request.clear_amount)
                    if requested_total > queue_authorities[0]:
                        raise ValueError("The requested amount exceeds the selected queue transaction value and charges.")
                if requested_total > available:
                    raise ValueError(
                        f"The requested {requested_total} exceeds the {request.currency} unauthorized debit of {available}."
                    )

            for index, (queue, queue_authority) in enumerate(zip(selected_queues, queue_authorities), start=1):
                if request.mode == "QUEUE_ROW_ONLY":
                    clear_amount = Decimal("0")
                    operation_mode = "QUEUE_ROW_ONLY"
                elif len(selected_queues) > 1:
                    clear_amount = queue_authority
                    operation_mode = "QUEUE_EXACT"
                else:
                    clear_amount = requested_total
                    operation_mode = "QUEUE_EXACT" if clear_amount == queue_authority else "QUEUE_PARTIAL"
                transactions.append(
                    ClearingTransactionInput(
                        source_row=index,
                        entry_date=_oracle_day(queue.get("entry_date")) or now.date(),
                        from_account=account,
                        source_internal_account=source_internal_account,
                        to_account="",
                        currency=request.currency,
                        rrn=_normalized_reference(queue.get("rrn"), "RRN"),
                        stan=_normalized_reference(queue.get("stan"), "STAN"),
                        amount=_money(queue.get("amount", 0)),
                        charge=_money(queue.get("fee", 0)),
                        clear_amount=clear_amount,
                        operation_mode=operation_mode,
                        queue_row_id=_normalized_text(queue.get("row_id")),
                        narration=request.note or (
                            "Account Explorer queue-only custody repair"
                            if operation_mode == "QUEUE_ROW_ONLY"
                            else "Account Explorer authorized debit clearing"
                        ),
                    )
                )

        primary_operation_mode = transactions[0].operation_mode
        scope_label = {
            "BALANCE_ALL": "full debit",
            "QUEUE_ROW_ONLY": (
                f"{len(transactions)} queue rows" if len(transactions) > 1 else f"queue row {transactions[0].rrn}"
            ),
            "QUEUE_AMOUNT_RESET": "queue amount repair",
        }.get(
            primary_operation_mode,
            f"{len(transactions)} queue rows" if len(transactions) > 1 else transactions[0].rrn,
        )
        batch = self.repository.create_batch(
            ClearingBatchCreateRequest(
                batch_name=f"{account} / {request.currency} / {scope_label}",
                finance_reference=request.change_reference,
                source_filename=source_filename,
                source_kind="ACCOUNT_EXPLORER",
                transactions=transactions,
            ),
            actor=actor,
            actor_role=actor_role,
        )
        return batch

    def create_account_batch(
        self,
        external_account: str,
        request: ClearingAccountBatchRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        if request.mode in {"QUEUE_TRANSACTION", "QUEUE_ROW_ONLY"}:
            self._assert_specific_record_allowed(actor_role=actor_role)
        account = _normalized_account(external_account)
        return self._create_account_batch_from_evidence(
            account,
            self.oracle.inspect_account(account),
            request,
            actor=actor,
            actor_role=actor_role,
        )

    def create_rrn_batch(
        self,
        rrn: str,
        request: ClearingAccountBatchRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        self._assert_specific_record_allowed(actor_role=actor_role)
        lookup_rrn = _normalized_rrn_lookup(rrn)
        evidence = self.oracle.inspect_rrn(lookup_rrn)
        internal_accounts = [str(value) for value in evidence.get("internal_accounts", [])]
        if len(internal_accounts) != 1:
            raise ValueError("The RRN must resolve to exactly one internal account before clearing can continue.")
        internal_account = internal_accounts[0]
        return self._create_account_batch_from_evidence(
            internal_account,
            evidence,
            request,
            actor=actor,
            actor_role=actor_role,
            source_internal_account=internal_account,
            source_filename=f"rrn-{lookup_rrn}.json",
        )

    def reconcile(self, batch_id: str, *, actor: str, actor_role: str) -> dict[str, Any]:
        self._assert_batch_window_open(batch_id)
        reconciliation_id = self.repository.begin_reconciliation(batch_id, actor)
        try:
            transactions = self.repository.selected_transactions(batch_id)
            external_accounts = {item["from_account"] for item in transactions}
            direct_mappings: dict[str, str] = {}
            for item in transactions:
                internal = _normalized_text(item.get("source_internal_account"))
                if not internal:
                    continue
                locator = item["from_account"]
                if locator in direct_mappings and direct_mappings[locator] != internal:
                    raise ValueError("One clearing locator resolves to conflicting internal accounts.")
                direct_mappings[locator] = internal
            collect_args = (
                external_accounts,
                {item["from_account"] for item in transactions},
            )
            oracle_state = (
                self.oracle.collect(*collect_args, direct_mappings=direct_mappings)
                if direct_mappings
                else self.oracle.collect(*collect_args)
            )
            results, snapshots, summary = reconcile_transactions(transactions, oracle_state)
            batch = self.repository.finish_reconciliation(
                batch_id,
                reconciliation_id,
                results=results,
                snapshots=snapshots,
                summary=summary,
                actor=actor,
                actor_role=actor_role,
            )
            audit_logger.log(
                event_type="nexus_unauthorized_clearing_reconciliation",
                user=actor,
                details={
                    "batch_id": batch_id,
                    "reconciliation_id": reconciliation_id,
                    "summary": summary,
                },
                success=True,
            )
            return batch
        except Exception as exc:
            self.repository.fail_reconciliation(batch_id, reconciliation_id, str(exc))
            audit_logger.log(
                event_type="nexus_unauthorized_clearing_reconciliation",
                user=actor,
                details={"batch_id": batch_id, "reconciliation_id": reconciliation_id, "error": str(exc)},
                success=False,
            )
            raise

    def submit(
        self,
        batch_id: str,
        request: ClearingSubmitRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        batch = self.repository.get_batch(batch_id)
        payload = build_execution_payload(batch, request.change_reference, request.note)
        return self.repository.save_submission(
            batch_id,
            payload=payload,
            payload_hash=payload["payload_hash"],
            actor=actor,
            actor_role=actor_role,
            change_reference=request.change_reference,
            note=request.note,
        )

    def approve(
        self,
        batch_id: str,
        request: ClearingApprovalRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        allow_self_approval = self._is_admin_role(actor_role)
        if not request.approve:
            return self.repository.decide_approval(
                batch_id,
                approve=False,
                actor=actor,
                actor_role=actor_role,
                note=request.note,
                allow_self_approval=allow_self_approval,
            )
        self._assert_batch_window_open(batch_id)
        idempotency_key = request.idempotency_key or f"approval:{batch_id}:{uuid4()}"
        with self.repository.action_lock(f"nexus-clearing:{batch_id}"):
            run = self.repository.approve_and_start_execution(
                batch_id,
                actor=actor,
                actor_role=actor_role,
                note=request.note.strip(),
                idempotency_key=idempotency_key,
                allow_self_approval=allow_self_approval,
            )
            return self._complete_execution(run, actor=actor, actor_role=actor_role)

    def _complete_execution(
        self,
        run: dict[str, Any],
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        if run["status"] in {"COMMITTED", "ROLLED_BACK", "COMMIT_UNCERTAIN"}:
            return run
        execution_id = run["execution_id"]
        payload = run["approved_payload"]
        if payload.get("payload_hash") != run["payload_hash"]:
            return self.repository.finish_execution(
                execution_id,
                status="BLOCKED",
                evidence={},
                error="Approved payload hash does not match the execution ledger.",
                actor=actor,
                actor_role=actor_role,
            )
        try:
            evidence = self.oracle.execute_atomic(
                payload,
                execution_id=execution_id,
                status_callback=lambda status: self.repository.update_execution_status(execution_id, status),
            )
            return self.repository.finish_execution(
                execution_id,
                status="COMMITTED",
                evidence=evidence,
                error=None,
                actor=actor,
                actor_role=actor_role,
            )
        except CommitUncertainError as exc:
            return self.repository.finish_execution(
                execution_id,
                status="COMMIT_UNCERTAIN",
                evidence={},
                error=str(exc),
                actor=actor,
                actor_role=actor_role,
            )
        except Exception as exc:
            logger.exception("Unauthorized clearing execution failed for %s", run.get("batch_id"))
            return self.repository.finish_execution(
                execution_id,
                status="BLOCKED",
                evidence={},
                error=str(exc),
                actor=actor,
                actor_role=actor_role,
            )

    def execution_preview(self, batch_id: str) -> dict[str, Any]:
        batch = self.repository.get_batch(batch_id)
        payload = batch.get("approved_payload")
        if not payload:
            raise ValueError("Submit the reconciled selection to create an immutable execution preview.")
        return {
            "batch_id": batch_id,
            "status": batch["status"],
            "payload_hash": batch["payload_hash"],
            "submitted_by": batch["submitted_by"],
            "approved_by": batch["approved_by"],
            "approved_at": batch["approved_at"],
            **payload,
        }

    def execute(
        self,
        batch_id: str,
        request: ClearingExecutionRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        self._assert_batch_window_open(batch_id)
        idempotency_key = request.idempotency_key or f"{batch_id}:{uuid4()}"
        with self.repository.action_lock(f"nexus-clearing:{batch_id}"):
            run = self.repository.start_execution(
                batch_id,
                actor=actor,
                actor_role=actor_role,
                reason=request.reason,
                idempotency_key=idempotency_key,
            )
            return self._complete_execution(run, actor=actor, actor_role=actor_role)

    def rollback(
        self,
        execution_id: str,
        request: ClearingRollbackRequest,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        with self.repository.action_lock(f"nexus-clearing-rollback:{execution_id}"):
            execution = self.repository.get_execution(execution_id)
            if execution["status"] != "COMMITTED":
                raise ValueError("Only a committed clearing execution can be compensated.")
            try:
                result = self.oracle.rollback_atomic(
                    execution["oracle_evidence"],
                    execution_id=execution_id,
                )
                return self.repository.finish_rollback(
                    execution_id,
                    actor=actor,
                    actor_role=actor_role,
                    reason=request.reason,
                    evidence=result,
                )
            except CommitUncertainError as exc:
                self.repository.record_rollback_failure(
                    execution_id,
                    actor=actor,
                    actor_role=actor_role,
                    reason=request.reason,
                    error=str(exc),
                    commit_uncertain=True,
                )
                raise
            except Exception as exc:
                self.repository.record_rollback_failure(
                    execution_id,
                    actor=actor,
                    actor_role=actor_role,
                    reason=request.reason,
                    error=str(exc),
                    commit_uncertain=False,
                )
                raise

    def account_view(self, batch_id: str, external_account: str) -> dict[str, Any]:
        batch = self.repository.get_batch(batch_id)
        account = _normalized_account(external_account)
        transactions = [
            item
            for item in batch["transactions"]
            if item["from_account"] == account
        ]
        if not transactions:
            raise LookupError("This account is not part of the batch debit scope.")
        latest_snapshot = None
        with self.repository._connection() as connection:
            with connection.cursor() as cursor:
                latest_snapshot = cursor.execute(
                    """
                    SELECT *
                    FROM nexus_clearing_account_snapshot
                    WHERE batch_id = %s AND external_account = %s
                    ORDER BY captured_at DESC
                    LIMIT 1
                    """,
                    (batch_id, account),
                ).fetchone()
        snapshot = _plain(dict(latest_snapshot)) if latest_snapshot else None
        if snapshot:
            snapshot["summary"] = summarize_account_evidence(
                snapshot.get("balance_rows", []),
                snapshot.get("queue_rows", []),
            )
            snapshot["internal_accounts"] = (
                [str(snapshot["internal_account"])]
                if snapshot.get("internal_account")
                else []
            )
            snapshot["source"] = "BATCH_SNAPSHOT"
        return {
            "external_account": account,
            "batch_id": batch_id,
            "source": "BATCH_SNAPSHOT",
            "transactions": transactions,
            "snapshot": snapshot,
        }

    def live_account_view(
        self,
        external_account: str,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        evidence = self.oracle.inspect_account(external_account)
        self.repository.record_account_lookup(
            evidence["external_account"],
            actor=actor,
            actor_role=actor_role,
            evidence_summary={
                "mapping_count": evidence["mapping_count"],
                **evidence["summary"],
            },
        )
        internal_accounts = evidence["internal_accounts"]
        return {
            "external_account": evidence["external_account"],
            "batch_id": None,
            "source": "LIVE_ORACLE",
            "lookup_mode": "ACCOUNT",
            "lookup_value": evidence["external_account"],
            "transactions": [],
            "snapshot": {
                "snapshot_id": None,
                "internal_account": internal_accounts[0] if len(internal_accounts) == 1 else None,
                "internal_accounts": internal_accounts,
                "account_role": "LIVE",
                "mapping_count": evidence["mapping_count"],
                "balance_rows": evidence["balance_rows"],
                "queue_rows": evidence["queue_rows"],
                "account_profiles": evidence.get("account_profiles", []),
                "resolution": {},
                "summary": evidence["summary"],
                "captured_at": evidence["captured_at"],
                "source": "LIVE_ORACLE",
            },
        }

    def live_rrn_view(
        self,
        rrn: str,
        *,
        actor: str,
        actor_role: str,
    ) -> dict[str, Any]:
        evidence = self.oracle.inspect_rrn(rrn)
        self.repository.record_rrn_lookup(
            evidence["lookup_value"],
            actor=actor,
            actor_role=actor_role,
            evidence_summary={
                "mapping_count": evidence["mapping_count"],
                "matched_queue_count": evidence["matched_queue_count"],
                "internal_accounts": evidence["internal_accounts"],
                "external_accounts": evidence.get("external_accounts", []),
                **evidence["summary"],
            },
        )
        internal_accounts = evidence["internal_accounts"]
        return {
            "external_account": evidence["external_account"],
            "batch_id": None,
            "source": "LIVE_ORACLE",
            "lookup_mode": "RRN",
            "lookup_value": evidence["lookup_value"],
            "transactions": [],
            "snapshot": {
                "snapshot_id": None,
                "internal_account": internal_accounts[0] if len(internal_accounts) == 1 else None,
                "internal_accounts": internal_accounts,
                "external_accounts": evidence.get("external_accounts", []),
                "account_role": "LIVE",
                "mapping_count": evidence["mapping_count"],
                "matched_queue_count": evidence["matched_queue_count"],
                "balance_rows": evidence["balance_rows"],
                "queue_rows": evidence["queue_rows"],
                "account_profiles": [],
                "resolution": {},
                "summary": evidence["summary"],
                "captured_at": evidence["captured_at"],
                "source": "LIVE_ORACLE",
            },
        }
